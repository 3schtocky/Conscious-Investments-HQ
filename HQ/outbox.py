"""The Outbox: what Client Relations prepares for the outside world, waiting for the Captain.

Nothing here publishes anything. A newsletter issue is a folder of ready-to-paste files:

    outbox/newsletters/<date>-<slug>/
        draft.md      what Wren and Harbor wrote (Markdown)
        issue.md      the final text with the disclaimer
        issue.html    the same as a web page, to open and read
        body.html     just the article, to paste into the Substack editor
        header.png    a branded header image
        social.md     one X post and one LinkedIn post
        meta.json     status, check results, approval card

    outbox/deliverables/<TICKER>/CI_<TICKER>_Memo_<date>.docx (+ .pdf when Word is available)

`check_issue` is the code gate (free): only Stott-approved numbers, no hype, no advice. An issue
with errors can't be finalized, so it never reaches the Captain's desk in that state.
Everything in outbox/ is gitignored.
"""

from __future__ import annotations

import html
import json
import re
import time
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING

from HQ.config import ROOT
from HQ.config import office as office_config

if TYPE_CHECKING:
    from HQ.engine.runtime import Office

OUTBOX_DIR = ROOT / "outbox"
ISSUE_ID = re.compile(r"^\d{4}-\d{2}-\d{2}-[a-z0-9][a-z0-9\-]{0,43}$")
WORDS_LOW, WORDS_HIGH = 350, 650       # "about 400 to 600 words"
X_LIMIT, LINKEDIN_LIMIT = 280, 1300
BODY_LIMIT = 12_000

DEFAULT_DISCLAIMER = (
    "Conscious Investments is an independent equity research imprint authored by Ethan Stott. It "
    "is not a registered investment adviser or broker-dealer. This note is for information and "
    "education only and is not investment advice, an offer to sell, or a solicitation to buy any "
    "security. Price targets and ratings reflect the author's judgment as of the date above, may "
    "change without notice and may not be realized. Past performance is not indicative of future "
    "results. Portions of this note were drafted with AI tools and reviewed by the author.")
DEFAULT_HOLDINGS = ("The author may hold positions in securities discussed. The Conscious "
                    "Investments model portfolio is a paper portfolio.")


def settings() -> dict:
    cfg = office_config().get("client_relations", {}) or {}
    return {"publisher": cfg.get("publisher", "outbox"),
            "disclaimer": " ".join((cfg.get("disclaimer") or DEFAULT_DISCLAIMER).split()),
            "holdings_disclosure": " ".join((cfg.get("holdings_disclosure") or DEFAULT_HOLDINGS).split())}


# ---- issues on disk ---------------------------------------------------------------------------
def slugify(title: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:40].strip("-")
    return slug or "issue"


def issue_dir(root: Path, issue_id: str) -> Path:
    if not ISSUE_ID.match(issue_id):
        raise KeyError(f"No newsletter issue {issue_id!r}. Use the id save_newsletter returned.")
    return root / "newsletters" / issue_id


def load(root: Path, issue_id: str) -> dict:
    path = issue_dir(root, issue_id) / "meta.json"
    if not path.is_file():
        raise KeyError(f"No newsletter issue {issue_id!r}. Use the id save_newsletter returned.")
    return json.loads(path.read_text())


def _save_meta(root: Path, meta: dict) -> dict:
    meta["updated"] = time.time()
    (issue_dir(root, meta["id"]) / "meta.json").write_text(json.dumps(meta, indent=1))
    return meta


def set_status(root: Path, issue_id: str, status: str, **extra) -> dict:
    meta = load(root, issue_id)
    meta.update(status=status, **extra)
    return _save_meta(root, meta)


def issues(root: Path) -> list[dict]:
    """Every issue, newest first."""
    out = []
    for path in (root / "newsletters").glob("*/meta.json"):
        try:
            out.append(json.loads(path.read_text()))
        except ValueError:
            continue
    return sorted(out, key=lambda m: m.get("updated", 0), reverse=True)


def deliverables(root: Path) -> list[dict]:
    out = []
    for path in sorted((root / "deliverables").glob("*/*.json")):
        try:
            out.append(json.loads(path.read_text()))
        except ValueError:
            continue
    return sorted(out, key=lambda m: m.get("updated", 0), reverse=True)


def word_count(body: str) -> int:
    text = re.sub(r"[#*_`|>\-]+", " ", re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", body))
    return len(re.findall(r"[A-Za-z0-9$%'’.]+", text))


LOCKED = {"awaiting": "waiting for a decision", "approved": "approved",
          "finalizing": "being finalized right now"}


def save_draft(office: Office, *, agent_id: str, title: str, body: str, x_post: str,
               linkedin_post: str, issue_id: str | None = None, day: str | None = None) -> dict:
    """Save a draft and run the code checks on it. With `issue_id` this revises that issue
    (whatever day it is and whatever the title now says); without it, the same title on the same
    day revises, and any other title starts a new issue."""
    root = office.outbox_dir
    if issue_id:
        old = load(root, issue_id)
    else:
        day = day or office.ledger.today()
        base = f"{day}-{slugify(title)}"
        issue_id, old, n = base, None, 1
        while (issue_dir(root, issue_id) / "meta.json").is_file():
            old = load(root, issue_id)
            if old["title"] == title:
                break
            n += 1                      # another issue already has this folder name
            issue_id, old = f"{base[:48]}-{n}", None
    if old and old["status"] in LOCKED:
        raise ValueError(f"Issue {issue_id} is {LOCKED[old['status']]}; it can't be changed. Use a "
                         "new title for a new issue.")
    folder = issue_dir(root, issue_id)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "draft.md").write_text(body.strip() + "\n")
    problems = check_issue(office, title=title, body=body, x_post=x_post, linkedin_post=linkedin_post)
    errors = [p for p in problems if p["level"] == "error"]
    old = old or {}
    meta = {"id": issue_id, "title": title, "date": old.get("date") or issue_id[:10],
            "status": "blocked" if errors else "draft",
            "words": word_count(body), "problems": problems, "x_post": x_post.strip(),
            "linkedin_post": linkedin_post.strip(), "drafted_by": old.get("drafted_by", agent_id),
            "revised_by": agent_id, "created": old.get("created", time.time()),
            "revisions": old.get("revisions", 0) + 1, "files": {}, "approval_id": None}
    return _save_meta(root, meta)


def recheck(office: Office, issue_id: str) -> list[dict]:
    """The checks again, against what is approved right now."""
    meta = load(office.outbox_dir, issue_id)
    body = (issue_dir(office.outbox_dir, issue_id) / "draft.md").read_text()
    return check_issue(office, title=meta["title"], body=body, x_post=meta["x_post"],
                       linkedin_post=meta["linkedin_post"])


# ---- the code gate ----------------------------------------------------------------------------
_RATING = re.compile(r"\b((?-i:Outperform|Underperform))\b|\b(?:we\s+rate|rating\s*(?::|of|is)|rated)\W+"
                     r"(?:\w+\W+){0,3}?(buy|sell|hold|neutral|overweight|underweight)\b", re.IGNORECASE)
_UPSIDE = re.compile(r"\d+(?:\.\d+)?\s?%\s+(?:upside|downside)|(?:upside|downside)\s+of\s+\d", re.IGNORECASE)
_EXPLICIT = re.compile(r"\$([A-Z]{1,5})\b|\((?:(?:NASDAQ|NYSE|AMEX)\s*:\s*)?([A-Z]{1,5})\)")
_HYPE = re.compile(r"guarantee\w*|can'?t\s+lose|risk[- ]free|sure\s+thing|to\s+the\s+moon|"
                   r"get\s+rich|once[- ]in[- ]a[- ]lifetime|no[- ]brainer|will\s+(?:double|triple|10x)|"
                   r"\bmoonshot\b|easy\s+money", re.IGNORECASE)
_ADVICE = re.compile(r"\byou\s+should\s+(?:buy|sell|own|hold|add|invest)\b|\bwe\s+recommend\s+(?:that\s+)?you\b|"
                     r"\b(?:buy|sell)\s+(?:it\s+)?now\b|\bbuy\s+this\s+stock\b|\byour\s+portfolio\s+should\b", re.IGNORECASE)
_EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿️]")
_NOT_TICKERS = {"A", "I", "AI", "US", "USA", "CEO", "CFO", "SEC", "ETF", "IPO", "GAAP", "EPS", "FY", "YOY",
                "PT", "SPY", "USD", "Q", "X", "IT", "ALL", "ON", "AT", "FOR", "ARE", "IS", "OR", "SO"}


def _problem(level: str, code: str, where: str, msg: str) -> dict:
    return {"level": level, "code": code, "where": where, "msg": msg}


def _usd(n: float) -> str:
    return f"${n:,.2f}"


def publishable(office: Office) -> dict:
    """What the office is allowed to put numbers on: approved models, and watchlist names
    (ideas only, no valuation)."""
    approved = {}
    for m in office.store.models():
        if m["status"] == "approved":
            pts = m["summary"].get("price_targets", {})
            approved[m["ticker"]] = {
                "version": m["version"], "rating": m["summary"].get("rating"),
                "targets": {k: float(v) for k, v in pts.items()
                            if isinstance(v, (int, float)) and not isinstance(v, bool)},
                "approved_at": m["approved_at"]}
    watch = {w["ticker"]: w for w in office.store.watchlist()}
    return {"approved": approved, "watch": watch}


def _check_text(where: str, text: str, pub: dict) -> list[dict]:
    from HQ import audit

    out: list[dict] = []
    approved, watch = pub["approved"], pub["watch"]
    known = set(approved) | set(watch)
    for sentence in audit.sentences(text):
        # A number belongs to the names in its own sentence: a target, rating or upside must sit
        # beside the ticker of an approved model, and beside no name that lacks one.
        named = {t for t in re.findall(r"\b[A-Z][A-Z.\-]{0,5}\b", sentence) if t in known}
        ok = sorted(t for t in named if t in approved)
        bare = sorted(named - set(ok))
        targets = [v for t in ok for v in approved[t]["targets"].values()]
        for n in ([] if audit.is_street(sentence) else audit.sentence_targets(sentence)):
            if bare:
                out.append(_problem("error", "unapproved-target", where,
                                    f"States a price target ({_usd(n)}) in a sentence about {', '.join(bare)}, "
                                    "which has no approved model. Watchlist names are ideas only."))
            elif not ok:
                out.append(_problem("error", "unapproved-target", where,
                                    f"Price target {_usd(n)} has no approved name beside it. Put the "
                                    "ticker of the approved model in the same sentence."))
            elif not audit.matches_target(n, targets):
                out.append(_problem("error", "unapproved-target", where,
                                    f"Price target {_usd(n)} does not match the approved model for "
                                    f"{', '.join(ok)}."))
        for m in _RATING.finditer(sentence):
            rating = (m.group(1) or m.group(2)).capitalize()
            allowed = {str(approved[t]["rating"]).capitalize() for t in ok}
            if bare:
                out.append(_problem("error", "unapproved-rating", where,
                                    f"States a rating ({rating}) in a sentence about {', '.join(bare)}, "
                                    "which has no approved model."))
            elif not ok:
                out.append(_problem("error", "unapproved-rating", where,
                                    f"Rating {rating} has no approved name beside it. Put the ticker "
                                    "of the approved model in the same sentence."))
            elif rating not in allowed:
                out.append(_problem("error", "unapproved-rating", where,
                                    f"Rating {rating} does not match the approved model "
                                    f"({', '.join(sorted(allowed))})."))
        if _UPSIDE.search(sentence) and (bare or not ok):
            out.append(_problem("error", "unapproved-target", where,
                                "States upside or downside with no approved model behind it."))
    for m in _EXPLICIT.finditer(text):
        t = m.group(1) or m.group(2)
        if t not in known and t not in _NOT_TICKERS:
            out.append(_problem("warning", "unknown-name", where,
                                f"{t} is neither a researched name nor on the watchlist."))
    if "—" in text or "–" in text and re.search(r"\s–\s", text):
        out.append(_problem("error", "em-dash", where, "No em dashes (the house style)."))
    for m in _HYPE.finditer(text):
        out.append(_problem("error", "hype", where, f"Hype: \"{m.group(0)}\". We never oversell."))
    for m in _ADVICE.finditer(text):
        out.append(_problem("error", "advice", where,
                            f"Reads as personal advice: \"{m.group(0)}\". This is research, not advice."))
    if "[VERIFY" in text:
        out.append(_problem("error", "verify", where, "A [VERIFY] marker is still in the text."))
    if _EMOJI.search(text):
        out.append(_problem("error", "emoji", where, "No emojis in published text."))
    try:
        from erb.lint import AI_ISMS
    except ImportError:   # pragma: no cover - erb is a dependency
        AI_ISMS = []
    low = text.lower()
    hits = sorted({w for w in AI_ISMS if re.search(r"\b" + re.escape(w) + r"\b", low)})
    if hits:
        out.append(_problem("warning", "ai-ism", where, f"Stock phrases to rewrite: {', '.join(hits)}."))
    seen, unique = set(), []
    for p in out:   # the same sentence can trip one rule twice
        key = (p["code"], p["msg"])
        if key not in seen:
            seen.add(key)
            unique.append(p)
    return unique


def check_issue(office: Office, *, title: str, body: str, x_post: str, linkedin_post: str) -> list[dict]:
    """Every reason this issue may not go out (errors) and things worth a second look (warnings)."""
    pub = publishable(office)
    out = _check_text("title", title, pub) + _check_text("newsletter", body, pub)
    words = word_count(body)
    if words < WORDS_LOW or words > WORDS_HIGH:
        out.append(_problem("warning", "length", "newsletter",
                            f"{words} words; the weekly note runs about 400 to 600."))
    for where, text, limit in (("X post", x_post, X_LIMIT), ("LinkedIn post", linkedin_post, LINKEDIN_LIMIT)):
        if not text.strip():
            out.append(_problem("error", "missing", where, f"The {where} is missing."))
            continue
        out += _check_text(where, text, pub)
        if len(text.strip()) > limit:
            out.append(_problem("error", "length", where,
                                f"{len(text.strip())} characters; the limit is {limit}."))
    return out


# ---- rendering --------------------------------------------------------------------------------
_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^\s)]+)\)")


def _inline(text: str) -> str:
    out = html.escape(text, quote=False)
    out = _LINK.sub(lambda m: f'<a href="{html.escape(html.unescape(m.group(2)), quote=True)}">{m.group(1)}</a>', out)
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    return re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"<em>\1</em>", out)


def markdown_to_html(md: str) -> str:
    """The small Markdown subset the newsletter uses: headings, paragraphs, bullet and numbered
    lists, bold, italic, links and pipe tables. Everything else is escaped as text."""
    out: list[str] = []
    for block in re.split(r"\n\s*\n", md.strip()):
        lines = [ln.rstrip() for ln in block.strip().splitlines() if ln.strip()]
        if not lines:
            continue
        head = re.match(r"^(#{1,3})\s+(.*)$", lines[0])
        if head and len(lines) == 1:
            level = len(head.group(1))
            out.append(f"<h{level}>{_inline(head.group(2))}</h{level}>")
        elif all(re.match(r"^\s*[-*]\s+", ln) for ln in lines):
            items = "".join(f"<li>{_inline(re.sub(r'^\s*[-*]\s+', '', ln))}</li>" for ln in lines)
            out.append(f"<ul>{items}</ul>")
        elif all(re.match(r"^\s*\d+[.)]\s+", ln) for ln in lines):
            items = "".join(f"<li>{_inline(re.sub(r'^\s*\d+[.)]\s+', '', ln))}</li>" for ln in lines)
            out.append(f"<ol>{items}</ol>")
        elif all(ln.lstrip().startswith("|") for ln in lines):
            rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
            if rows:
                thead = "".join(f"<th>{_inline(c)}</th>" for c in rows[0])
                tbody = "".join("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in r) + "</tr>"
                                for r in rows[1:])
                out.append(f"<table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table>")
        else:
            if head:   # a heading followed directly by text
                level = len(head.group(1))
                out.append(f"<h{level}>{_inline(head.group(2))}</h{level}>")
                lines = lines[1:]
            if lines:
                out.append(f"<p>{_inline(' '.join(ln.strip() for ln in lines))}</p>")
    return "\n".join(out)


PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  body {{ margin: 0; background: #ECEAE1; color: #2E2C29; font: 18px/1.6 Georgia, serif; }}
  article {{ max-width: 680px; margin: 0 auto; padding: 24px 20px 56px; background: #fff; }}
  img {{ width: 100%; height: auto; border-radius: 6px; }}
  h1 {{ font-size: 30px; line-height: 1.2; margin: 20px 0 4px; }}
  h2 {{ font-size: 21px; margin: 28px 0 6px; }} h3 {{ font-size: 18px; margin: 20px 0 4px; }}
  .date {{ color: #5F5B55; font: 14px/1.4 system-ui, sans-serif; margin: 0 0 18px; }}
  table {{ border-collapse: collapse; width: 100%; font: 15px/1.4 system-ui, sans-serif; }}
  th, td {{ border-bottom: 1px solid #E3E0D8; padding: 6px 8px; text-align: left; }}
  .disclaimer {{ color: #5F5B55; font: 13px/1.5 system-ui, sans-serif; border-top: 1px solid #E3E0D8; margin-top: 32px; padding-top: 12px; }}
  a {{ color: #2a68a8; }}
</style></head><body><article>
<img src="header.png" alt="">
<h1>{title}</h1>
<p class="date">{date}</p>
{body}
</article></body></html>
"""


def _font(size: int, bold: bool = False):
    from PIL import ImageFont

    names = (["Georgia Bold.ttf", "Georgia.ttf"] if bold else ["Georgia.ttf"]) + ["DejaVuSerif.ttf"]
    for name in names:
        for base in ("/System/Library/Fonts/Supplemental/", "/Library/Fonts/", ""):
            try:
                return ImageFont.truetype(base + name, size)
            except OSError:
                continue
    return ImageFont.load_default(size)


def header_image(title: str, date_text: str, path: Path) -> Path:
    """A branded header (1456 x 816): paper background, the logo, the title, an accent rule."""
    from erb.config import ASSETS_DIR
    from PIL import Image, ImageDraw

    w, h = 1456, 816
    img = Image.new("RGB", (w, h), "#ECEAE1")
    draw = ImageDraw.Draw(img)
    for i, colour in enumerate(("#2a68a8", "#d27030", "#3a9a7a", "#8a6db8")):   # the brand order
        draw.rectangle([i * w // 4, h - 18, (i + 1) * w // 4, h], fill=colour)
    x, y = 96, 96
    try:
        logo = Image.open(ASSETS_DIR / "logo_transparent.png").convert("RGBA")
        logo.thumbnail((360, 150))
        img.paste(logo, (x, y), logo)
        y += logo.height + 56
    except OSError:
        y += 40
    draw.text((x, y), "CONSCIOUS INVESTMENTS  ·  WEEKLY NOTE", font=_font(30), fill="#5F5B55")
    y += 70
    font = _font(84, bold=True)
    words, lines, line = title.split(), [], ""
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textlength(trial, font=font) <= w - 2 * x or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    lines.append(line)
    for ln in lines[:3]:
        draw.text((x, y), ln, font=font, fill="#2E2C29")
        y += 104
    draw.text((x, h - 110), date_text, font=_font(32), fill="#5F5B55")
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=True)
    return path


def _pretty(day: str) -> str:
    d = date.fromisoformat(day)
    return f"{d.strftime('%B')} {d.day}, {d.year}"


def render_issue(office: Office, issue_id: str) -> dict:
    """Write the final files for an issue that passed the checks. Returns the updated meta."""
    root = office.outbox_dir
    meta = load(root, issue_id)
    folder = issue_dir(root, issue_id)
    cfg = settings()
    body = (folder / "draft.md").read_text().strip()
    pretty = _pretty(meta["date"])
    disclaimer = f"{cfg['disclaimer']} Holdings disclosure: {cfg['holdings_disclosure']}"
    (folder / "issue.md").write_text(f"# {meta['title']}\n\n{pretty}\n\n{body}\n\n---\n\n*{disclaimer}*\n")
    article = markdown_to_html(body) + f'\n<p class="disclaimer"><em>{html.escape(disclaimer)}</em></p>'
    (folder / "body.html").write_text(article + "\n")
    (folder / "issue.html").write_text(PAGE.format(title=html.escape(meta["title"]), date=pretty, body=article))
    header_image(meta["title"], pretty, folder / "header.png")
    (folder / "social.md").write_text(f"# X\n\n{meta['x_post']}\n\n# LinkedIn\n\n{meta['linkedin_post']}\n")
    meta["files"] = {k: f"newsletters/{issue_id}/{k}" for k in
                     ("issue.md", "issue.html", "body.html", "header.png", "social.md")}
    return _save_meta(root, meta)


# ---- client-ready memos -----------------------------------------------------------------------
def package_memo(office: Office, ticker: str, *, pdf: bool = True) -> dict:
    """Turn a finished research memo into a branded client file (.docx, plus .pdf when Word is
    available). The memo must use the approved model's numbers and have nothing left to verify."""
    from HQ import audit
    from HQ.tools import desk

    cdir = desk.coverage_dir(ticker)
    memo = cdir / "memo.md"
    if not memo.is_file():
        raise ValueError(f"No memo.md for {ticker} yet; Research writes it once the model is approved.")
    model = office.store.approved_model(ticker)
    if model is None:
        raise ValueError(f"{ticker} has no approved model, so there are no numbers to send to a client.")
    text = memo.read_text(errors="replace")
    found = audit.check_document("memo", text, subject=f"{ticker} memo.md", ticker=ticker,
                                 approved={"version": model["version"],
                                           "targets": list(publishable(office)["approved"][ticker]["targets"].values())},
                                 has_sources=(cdir / "sources.md").is_file())
    problems = [f.detail for f in found] + [p["msg"] for p in _check_text("memo", text, publishable(office))
                                            if p["level"] == "error" and p["code"] in ("hype", "advice", "emoji")]
    if problems:
        raise ValueError("The memo isn't client-ready:\n- " + "\n- ".join(problems))

    day = office.ledger.today()
    current = deliverable(office.outbox_dir, f"{ticker}-{day}")
    if current and current["status"] in ("awaiting", "approved"):
        what = (f"waiting for a decision (approval #{current.get('approval_id')})"
                if current["status"] == "awaiting" else "already approved")
        raise ValueError(f"Today's {ticker} client memo is {what}; it can't be replaced. If the memo "
                         "changed, ask for the card to be sent back first.")
    out_dir = office.outbox_dir / "deliverables" / ticker
    out_dir.mkdir(parents=True, exist_ok=True)
    docx = out_dir / f"CI_{ticker}_Memo_{day}.docx"
    approved_on = datetime.fromtimestamp(model["approved_at"], office.ledger.tz).date().isoformat() \
        if model["approved_at"] else day
    _memo_docx(text, docx, ticker=ticker, day=day,
               model_line=f"Figures from Quant model v{model['version']}, approved {approved_on}.",
               holdings=settings()["holdings_disclosure"])
    files = {"docx": f"deliverables/{ticker}/{docx.name}"}
    if pdf:
        from erb import word

        made = word.finalize(docx, log=lambda *_: None)
        if made is not None and Path(made).is_file():
            files["pdf"] = f"deliverables/{ticker}/{Path(made).name}"
    meta = {"id": f"{ticker}-{day}", "ticker": ticker, "date": day, "title": f"{ticker} research memo",
            "status": "awaiting", "files": files, "model_version": model["version"],
            "approval_id": None, "updated": time.time()}
    (out_dir / f"{day}.json").write_text(json.dumps(meta, indent=1))
    return meta


def _deliverable_path(root: Path, deliverable_id: str) -> Path | None:
    from HQ.tools import desk

    m = re.fullmatch(r"(.+)-(\d{4}-\d{2}-\d{2})", deliverable_id)
    if not m or not desk.TICKER.match(m.group(1)):
        return None
    path = root / "deliverables" / m.group(1) / f"{m.group(2)}.json"
    return path if path.is_file() else None


def deliverable(root: Path, deliverable_id: str) -> dict | None:
    path = _deliverable_path(root, deliverable_id)
    return json.loads(path.read_text()) if path else None


def set_deliverable_status(root: Path, deliverable_id: str, status: str, **extra) -> dict:
    path = _deliverable_path(root, deliverable_id)
    if path is None:
        raise KeyError(f"No client deliverable {deliverable_id!r} in the Outbox.")
    meta = json.loads(path.read_text())
    meta.update(status=status, updated=time.time(), **extra)
    path.write_text(json.dumps(meta, indent=1))
    return meta


def _memo_docx(text: str, out: Path, *, ticker: str, day: str, model_line: str, holdings: str) -> None:
    from docx import Document
    from docx.shared import Inches
    from erb import docx_build as db
    from erb.config import ASSETS_DIR

    doc = Document()
    db.setup_styles(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(db.MARGIN)
    sec.top_margin, sec.bottom_margin = Inches(0.55), Inches(0.6)
    logo = ASSETS_DIR / db.BRAND.get("logo", "logo_transparent.png")
    if logo.is_file():
        doc.add_picture(str(logo), width=Inches(1.6))
    pretty = _pretty(day)
    db.small(doc, f"Conscious Investments  ·  Research memo  ·  {pretty}", size=8)
    db.small(doc, model_line, size=8, italic=True)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s or s.startswith("<!--"):
            i += 1
        elif s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            db.md_table(doc, block)
        elif m := re.match(r"^(#{1,3})\s+(.*)$", s):
            doc.add_heading(db.SOURCE_TAG.sub("", m.group(2)), level=len(m.group(1)))
            i += 1
        elif re.match(r"^[-*]\s+", s):
            db.add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", s))
            i += 1
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3}\s|\||[-*]\s)", lines[i].strip()):
                para.append(lines[i].strip())
                i += 1
            db.body_para(doc, " ".join(para))
    db.disclaimer(doc, holdings)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
