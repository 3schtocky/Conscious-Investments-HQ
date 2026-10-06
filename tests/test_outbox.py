"""Phase 7: Client Relations. The newsletter's code gate, the Outbox, the Publisher and client
memo packaging. Nothing is ever posted; everything waits for the Captain."""

from __future__ import annotations

import json

import pytest
from conftest import nick, register_model, text_turn, tool_turn

from hq import outbox
from hq.tools import desk

BODY = ("## The idea\n\nRambus is the name we finished researching this week. Our base price target "
        "for RMBS is $2.00 and we rate it Neutral, per Quant model v1.\n\n## What we're watching\n\n"
        "- **BFLY**: revenue growth is speeding up and margins are widening.\n")
X = "This week: RMBS at a Neutral rating, and BFLY on the watchlist as an idea."
LI = "Our weekly note covers RMBS and what we are watching."


def _results(llm, who, call=1):
    return [c for c in llm.calls if c["who"] == who][call]["params"]["messages"][-1]["content"]


@pytest.fixture
def office_llm(make_office, monkeypatch):
    """An office with one approved model (RMBS: 1 / 2 / 3, Neutral) and one watchlist name."""
    office, llm = make_office()
    register_model(office, "RMBS", 1)
    office.store.set_model_status("RMBS", 1, "approved")
    office.add_watch(ticker="BFLY", added_by="screen_lead", source="gems 2026-09-30 #1",
                     thesis="Growth is speeding up", pitch=None, price=10.0, spy=500.0)
    monkeypatch.setattr("hq.quotes.latest", lambda tickers: {t: {"BFLY": 11.0, "SPY": 505.0}.get(t) for t in tickers})
    return office, llm


def codes(office, body=BODY, x=X, li=LI, title="Weekly note"):
    return {(p["code"], p["level"], p["where"])
            for p in outbox.check_issue(office, title=title, body=body, x_post=x, linkedin_post=li)}


# ---- the code gate --------------------------------------------------------------------------
def test_only_approved_numbers_may_be_published(office_llm):
    office, _ = office_llm
    errors = lambda **kw: {c for c in codes(office, **kw) if c[1] == "error"}
    assert errors() == set()
    # a target that isn't the model's, or a rating that isn't
    assert ("unapproved-target", "error", "newsletter") in errors(body=BODY.replace("$2.00", "$2.40"))
    assert ("unapproved-rating", "error", "newsletter") in errors(body=BODY.replace("Neutral", "Outperform"))
    # watchlist names are ideas only: no target, rating or upside
    watch = "## Watching\n\nBFLY: our price target is $20.\n"
    assert ("unapproved-target", "error", "newsletter") in errors(body=watch)
    assert ("unapproved-target", "error", "newsletter") in errors(body="BFLY has 40% upside from here.")
    assert ("unapproved-rating", "error", "newsletter") in errors(body="We rate BFLY a buy.")
    # a number needs its approved ticker in the same sentence, and can't borrow a neighbour's
    assert ("unapproved-target", "error", "newsletter") in errors(body="RMBS is solid. Our price target is $3.00.")
    assert ("unapproved-target", "error", "newsletter") in errors(body="RMBS looks solid. BFLY: price target $2.00.")
    assert ("unapproved-rating", "error", "newsletter") in errors(body="RMBS is fine. BFLY is rated Neutral by us.")
    assert ("unapproved-target", "error", "newsletter") in errors(body="RMBS and BFLY: price target $2.00.")
    # a word like "analyst" or "average" does not excuse our own claim
    sneaky = "Our analyst team puts a price target of $500 on BFLY, 80% upside, rated Buy."
    assert {c[0] for c in errors(body=sneaky)} == {"unapproved-target", "unapproved-rating"}
    assert ("unapproved-target", "error", "newsletter") in errors(body="On average we see a price target of $9 for RMBS.")
    # the Street's view with attribution, and ordinary English, are fine
    assert errors(body="Consensus price target for BFLY is $14 [F3]. We expect it to outperform peers.") == set()
    # the social posts are held to the same rule
    assert ("unapproved-target", "error", "X post") in errors(x="BFLY price target $20!")


def test_house_rules_on_voice_and_format(office_llm):
    office, _ = office_llm
    got = lambda **kw: codes(office, **kw)
    assert ("hype", "error", "newsletter") in got(body="This one is a guaranteed winner.")
    assert ("advice", "error", "newsletter") in got(body="You should buy RMBS today.")
    assert ("em-dash", "error", "newsletter") in got(body="Growth — real growth — is back.")
    assert ("emoji", "error", "LinkedIn post") in got(li="Our note is out \U0001F680")
    assert ("verify", "error", "newsletter") in got(body="Revenue grew [VERIFY: figure].")
    assert ("length", "error", "X post") in got(x="x" * 281)
    assert ("missing", "error", "LinkedIn post") in got(li="  ")
    assert ("hype", "error", "title") in got(title="A sure thing in chips")
    assert ("length", "warning", "newsletter") in got()                      # the short test body
    assert ("ai-ism", "warning", "newsletter") in got(body="We delve into a robust quarter.")
    assert ("unknown-name", "warning", "newsletter") in got(body="We also looked at Zeta (ZETA).")
    ok = " ".join(["Plain words about how the screen works and why patience pays."] * 45)
    assert {c for c in got(body=ok) if c[0] == "length"} == set()


def test_markdown_becomes_safe_html():
    html = outbox.markdown_to_html(
        "## Idea <script>alert(1)</script>\n\nA **bold** and *plain* [link](https://example.com/a?b=1&c=2) "
        "and a [bad](javascript:alert(1)).\n\n- one\n- two\n\n1. first\n2. second\n\n"
        "| Name | Return |\n|---|---|\n| BFLY | +3.0% |")
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "<strong>bold</strong>" in html and "<em>plain</em>" in html
    assert '<a href="https://example.com/a?b=1&amp;c=2">link</a>' in html
    assert "javascript:" not in html.replace("[bad](javascript:alert(1))", "")   # left as plain text
    assert "<ul><li>one</li><li>two</li></ul>" in html and "<ol><li>first</li>" in html
    assert "<th>Name</th>" in html and "<td>+3.0%</td>" in html


# ---- the tools --------------------------------------------------------------------------------
def test_client_relations_tools_by_role():
    from hq.tools.office import tools_for

    names = lambda tier, aid, wing: {t.name for t in tools_for(tier, aid, wing)}
    harbor, wren = names("lead", "cr_lead", "client_relations"), names("associate", "cr_associate", "client_relations")
    assert {"newsletter_material", "save_newsletter", "read_newsletter", "finalize_newsletter", "package_memo"} <= harbor
    assert {"newsletter_material", "save_newsletter", "read_newsletter"} <= wren
    assert not {"finalize_newsletter", "package_memo"} & wren          # only the lead sends to the Captain
    assert not {"save_newsletter", "finalize_newsletter"} & names("lead", "er_lead", "equity_research")


async def test_newsletter_from_draft_to_ready_to_paste(office_llm):
    office, llm = office_llm
    bad = {"title": "Weekly note", "body": BODY.replace("$2.00", "$9.99"), "x_post": X, "linkedin_post": LI}
    good = {**bad, "body": BODY}
    issue = f"{office.ledger.today()}-weekly-note"
    llm.script("cr_lead", tool_turn(("delegate", {"to": "cr_associate", "job": "Draft the weekly note."})),
               tool_turn(("finalize_newsletter", {"issue": issue})),             # refused: still blocked
               tool_turn(("save_newsletter", good), ("finalize_newsletter", {"issue": issue, "note": "First issue."})),
               tool_turn(("save_newsletter", good), ("request_approval", {"kind": "newsletter", "title": "x", "summary": "y"})),
               text_turn("In the Outbox."), text_turn("Thanks."))
    llm.script("cr_associate", tool_turn(("newsletter_material", {})), tool_turn(("save_newsletter", bad)),
               tool_turn(("submit_result", {"findings": f"Saved as {issue}", "confidence": "medium"})))
    office.assign("cr_lead", "Prepare this week's newsletter")
    await office.idle()

    material = json.loads(_results(llm, "cr_associate", 1)[0]["content"])
    assert material["researched_names"] == [{"ticker": "RMBS", "rating": "Neutral",
                                             "price_targets": {"bear": 1.0, "base": 2.0, "bull": 3.0},
                                             "cite_as": f"Quant model v1, approved {office.ledger.today()}"}]
    [w] = material["watchlist"]
    assert (w["ticker"], w["return_since_flagged"], w["vs_sp500"]) == ("BFLY", 0.1, 0.09)
    saved = json.loads(_results(llm, "cr_associate", 2)[0]["content"])
    assert saved["status"] == "blocked" and "does not match the approved model for RMBS" in saved["errors"][0]

    refused = _results(llm, "cr_lead", 2)[0]
    assert refused["is_error"] and "can't go to Stott yet" in refused["content"]
    done = _results(llm, "cr_lead", 3)
    assert json.loads(done[0]["content"])["errors"] == [] and "approval #1 is on Stott's desk" in done[1]["content"]
    late = _results(llm, "cr_lead", 4)
    assert late[0]["is_error"] and "waiting for a decision" in late[0]["content"]       # frozen once filed
    assert late[1]["is_error"] and "finalize_newsletter" in late[1]["content"]          # no side door

    [card] = office.store.approvals("pending")
    assert card["kind"] == "newsletter" and card["payload"]["issue"] == issue
    assert f"/outbox/newsletters/{issue}/issue.html" in card["payload"]["attachments"]
    assert "Nothing is posted anywhere" in card["summary"] and "First issue." in card["summary"]
    folder = office.outbox_dir / "newsletters" / issue
    final = (folder / "issue.md").read_text()
    assert final.startswith("# Weekly note") and "price target for RMBS is $2.00" in final
    assert "not investment advice" in final and "Holdings disclosure: The author holds no position" in final
    page = (folder / "issue.html").read_text()
    assert "<h2>The idea</h2>" in page and 'class="disclaimer"' in page and '<img src="header.png"' in page
    assert (folder / "body.html").read_text().startswith("<h2>The idea</h2>")
    assert (folder / "header.png").read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"
    assert (folder / "social.md").read_text() == f"# X\n\n{X}\n\n# LinkedIn\n\n{LI}\n"
    meta = outbox.load(office.outbox_dir, issue)
    assert (meta["status"], meta["approval_id"], meta["drafted_by"], meta["revised_by"]) == \
        ("awaiting", card["id"], "cr_associate", "cr_lead")

    office.decide(card["id"], "approved", "Nice.")
    await office.idle()
    meta = outbox.load(office.outbox_dir, issue)
    assert meta["status"] == "approved" and "copy the article into Substack" in meta["next_step"]
    view = office.outbox_view()
    assert view["publisher"] == "outbox" and view["issues"][0]["drafted_by_name"] == nick("cr_associate")


async def test_changes_requested_reopens_the_issue(office_llm):
    office, llm = office_llm
    draft = {"title": "Weekly note", "body": BODY, "x_post": X, "linkedin_post": LI}
    issue = f"{office.ledger.today()}-weekly-note"
    llm.script("cr_lead", tool_turn(("save_newsletter", draft), ("finalize_newsletter", {"issue": issue})),
               text_turn("Filed."),
               tool_turn(("save_newsletter", {**draft, "issue": issue, "title": "Weekly note, revised",
                                              "body": BODY + "\nA sharper close.\n"}),
                         ("finalize_newsletter", {"issue": issue})),
               text_turn("Revised and filed again."))
    office.assign("cr_lead", "Newsletter")
    await office.idle()
    office.decide(office.store.approvals("pending")[0]["id"], "changes", "Sharper close, please.")
    await office.idle()
    meta = outbox.load(office.outbox_dir, issue)
    assert meta["status"] == "awaiting" and meta["revisions"] == 2 and meta["approval_id"] == 2
    assert meta["title"] == "Weekly note, revised" and len(outbox.issues(office.outbox_dir)) == 1
    assert "A sharper close." in (office.outbox_dir / "newsletters" / issue / "issue.md").read_text()


async def test_an_approval_withdrawn_after_drafting_blocks_the_issue(office_llm):
    office, llm = office_llm
    issue = f"{office.ledger.today()}-weekly-note"
    outbox.save_draft(office, agent_id="cr_associate", title="Weekly note", body=BODY, x_post=X, linkedin_post=LI)
    office.store.set_model_status("RMBS", 1, "superseded")      # the model stopped being official
    llm.script("cr_lead", tool_turn(("finalize_newsletter", {"issue": issue}), ("read_newsletter", {"issue": "nope"})),
               text_turn("Blocked."))
    office.assign("cr_lead", "Finalize")
    await office.idle()
    r = _results(llm, "cr_lead")
    assert r[0]["is_error"] and "no approved name beside it" in r[0]["content"]
    assert r[1]["is_error"] and "No newsletter issue" in r[1]["content"]
    assert outbox.load(office.outbox_dir, issue)["status"] == "blocked" and not office.store.approvals()


# ---- client memos -----------------------------------------------------------------------------
@pytest.fixture
def coverage(tmp_path, monkeypatch):
    root = tmp_path / "erb"
    (root / "coverage").mkdir(parents=True)
    monkeypatch.setattr(desk, "ERB_DIR", root)
    monkeypatch.setattr("erb.word.finalize", lambda docx, log=print: None)   # no Microsoft Word in tests
    return root / "coverage"


async def test_package_memo_only_ships_approved_finished_work(office_llm, coverage):
    from docx import Document

    office, llm = office_llm
    folder = coverage / "RMBS"
    folder.mkdir()
    calls = [("package_memo", {"ticker": "ZZZZ"}), ("package_memo", {"ticker": "RMBS"})]
    llm.script("cr_lead", tool_turn(*calls), text_turn("Not ready."))
    (folder / "memo.md").write_text("# Rambus\n\nBase price target $5.00 [M]. Growth was [VERIFY: 10-Q].\n")
    office.assign("cr_lead", "Package the memos")
    await office.idle()
    r = _results(llm, "cr_lead")
    assert r[0]["is_error"] and "No memo.md for ZZZZ" in r[0]["content"]
    assert r[1]["is_error"] and "isn't client-ready" in r[1]["content"]
    assert "[VERIFY]" in r[1]["content"] and "does not match approved model v1" in r[1]["content"]
    assert "no sources.md" in r[1]["content"] and not office.store.approvals()

    (folder / "memo.md").write_text("# Rambus: steady, fairly priced\n\nOur base price target is $2.00 [M] "
                                    "and we rate the shares Neutral.\n\n## Risks\n\n- Customer concentration [S1]\n\n"
                                    "| Scenario | Target |\n|---|---|\n| Bear | $1.00 |\n| Bull | $3.00 |\n")
    (folder / "sources.md").write_text("- [S1] 10-K FY'25\n")
    llm.script("cr_lead", tool_turn(("package_memo", {"ticker": "rmbs"})), text_turn("Packaged."), text_turn("Thanks."))
    office.assign("cr_lead", "Package RMBS")
    await office.idle()
    [card] = office.store.approvals("pending")
    day = office.ledger.today()
    assert card["kind"] == "deliverable" and card["payload"]["deliverable"] == f"RMBS-{day}"
    assert card["payload"]["attachments"] == [f"/outbox/deliverables/RMBS/CI_RMBS_Memo_{day}.docx"]
    doc = Document(str(office.outbox_dir / "deliverables" / "RMBS" / f"CI_RMBS_Memo_{day}.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Figures from Quant model v1" in text and "Our base price target is $2.00 and we rate" in text
    assert "[M]" not in text and "[S1]" not in text                        # source tags stay in-house
    assert "Holdings disclosure: The author holds no position" in text and len(doc.tables) == 1
    assert office.outbox_view()["deliverables"][0]["status"] == "awaiting"
    office.decide(card["id"], "approved")
    await office.idle()
    assert office.outbox_view()["deliverables"][0]["status"] == "approved"


# ---- endpoints --------------------------------------------------------------------------------
def test_outbox_endpoints(make_office):
    from fastapi.testclient import TestClient

    from hq.server import create_app

    office, _ = make_office()
    meta = outbox.save_draft(office, agent_id="cr_associate", title="Weekly note", body="Plain words.",
                             x_post="Out now.", linkedin_post="Out now.")
    outbox.render_issue(office, meta["id"])
    (office.outbox_dir.parent / "secret.txt").write_text("no")
    with TestClient(create_app(office_factory=lambda: office)) as c:
        view = c.get("/api/outbox").json()
        assert view["issues"][0]["id"] == meta["id"] and view["issues"][0]["status"] == "draft"
        assert c.get(f"/outbox/newsletters/{meta['id']}/issue.html").status_code == 200
        assert c.get(f"/outbox/newsletters/{meta['id']}/header.png").headers["content-type"] == "image/png"
        assert c.get("/outbox/%2e%2e/secret.txt").status_code == 404
        assert c.get("/outbox/newsletters/nope/issue.html").status_code == 404


def test_unknown_publisher_is_named(make_office):
    from hq.publish import OutboxPublisher, get_publisher

    assert isinstance(get_publisher(), OutboxPublisher) and get_publisher("outbox").name == "outbox"
    with pytest.raises(ValueError, match="Unknown publisher 'substack'"):
        get_publisher("substack")


# ---- the demo scene ---------------------------------------------------------------------------
async def test_demo_newsletter_scene_passes_the_real_checks(office_llm):
    from hq.demo import DemoLLM, scene_newsletter

    office, _ = office_llm
    llm = DemoLLM(speed=1000, office=office)
    office._llm = llm
    scene_newsletter(office, llm)
    await office.idle()
    assert all(t["status"] == "done" for t in office.store.tasks())
    [meta] = outbox.issues(office.outbox_dir)
    assert meta["status"] == "awaiting" and meta["problems"] == [], meta["problems"]
    assert outbox.WORDS_LOW <= meta["words"] <= outbox.WORDS_HIGH
    assert "**BFLY**" in (office.outbox_dir / "newsletters" / meta["id"] / "issue.md").read_text()
    assert len(meta["x_post"]) <= 280 and not any(llm.scripts.values())
    types = {e["type"] for e in office.store.events(limit=2000)}
    assert "guard_block" not in types and "task_error" not in types
    assert office.store.findings() == []                 # and Audit's own checks agree


# ---- regressions from the Phase 7 code review -------------------------------------------------
def test_issue_ids_never_collide_and_revisions_keep_their_id(office_llm):
    office, _ = office_llm
    save = lambda title, **kw: outbox.save_draft(office, agent_id="cr_associate", title=title, body="Plain words.",
                                                 x_post="Out.", linkedin_post="Out.", **kw)
    long = "Weekly note: growth that is speeding up in "
    a, b = save(long + "semis"), save(long + "software")
    assert a["id"] != b["id"] and len(outbox.issues(office.outbox_dir)) == 2
    assert save(long + "semis")["id"] == a["id"]                         # same title, same day: a revision
    later = save("A new title", issue_id=a["id"], day="2026-12-31")      # revised another day, by id
    assert later["id"] == a["id"] and later["revisions"] == 3 and later["date"] == a["date"]
    assert len(outbox.issues(office.outbox_dir)) == 2
    with pytest.raises(KeyError):
        save("x", issue_id="2026-01-01-missing")


async def test_approval_is_refused_when_the_approved_numbers_moved(office_llm):
    office, llm = office_llm
    issue = f"{office.ledger.today()}-weekly-note"
    llm.script("cr_lead", tool_turn(("save_newsletter", {"title": "Weekly note", "body": BODY, "x_post": X,
                                                        "linkedin_post": LI}),
                                   ("finalize_newsletter", {"issue": issue})), text_turn("Filed."), text_turn("Ok."))
    office.assign("cr_lead", "Newsletter")
    await office.idle()
    [card] = office.store.approvals("pending")
    office.store.set_model_status("RMBS", 1, "superseded")               # Quant's numbers moved on
    with pytest.raises(ValueError, match="changed since this issue was finalized"):
        office.decide(card["id"], "approved")
    assert office.store.approval(card["id"])["status"] == "pending"      # nothing was recorded
    assert outbox.load(office.outbox_dir, issue)["status"] == "awaiting"
    office.decide(card["id"], "changes", "Revise against the new model.")
    await office.idle()
    assert outbox.load(office.outbox_dir, issue)["status"] == "changes"


async def test_a_draft_cannot_change_while_it_is_being_finalized(office_llm):
    office, _ = office_llm
    meta = outbox.save_draft(office, agent_id="cr_associate", title="Weekly note", body=BODY, x_post=X, linkedin_post=LI)
    outbox.set_status(office.outbox_dir, meta["id"], "finalizing")
    with pytest.raises(ValueError, match="being finalized right now"):
        outbox.save_draft(office, agent_id="cr_associate", title="Weekly note", body="Price target $99.",
                          x_post=X, linkedin_post=LI)


async def test_a_packaged_memo_is_not_replaced_while_on_the_captains_desk(office_llm, coverage):
    office, llm = office_llm
    office.store.add_model(ticker="BRK-B", version=1, path="/tmp/b.xlsx", created_by="quant_lead",
                           summary={"price_targets": {"bear": 1, "base": 2, "bull": 3}, "rating": "Neutral",
                                    "check": {"ok": True}})
    office.store.set_model_status("BRK-B", 1, "approved")
    folder = coverage / "BRK-B"
    folder.mkdir()
    (folder / "memo.md").write_text("# Berkshire\n\nBRK-B base price target $2.00 [M].\n")
    (folder / "sources.md").write_text("- [S1] 10-K\n")
    llm.script("cr_lead", tool_turn(("package_memo", {"ticker": "BRK-B"})),
               tool_turn(("package_memo", {"ticker": "BRK-B"})), text_turn("Packaged once."), text_turn("Ok."))
    office.assign("cr_lead", "Package BRK-B")
    await office.idle()
    again = _results(llm, "cr_lead", 2)[0]
    assert again["is_error"] and "waiting for a decision (approval #1)" in again["content"]
    [card] = office.store.approvals("pending")
    office.decide(card["id"], "approved")                                # a hyphenated ticker's id still resolves
    await office.idle()
    assert office.outbox_view()["deliverables"][0]["status"] == "approved"


def test_newsletter_kind_is_not_offered_on_the_generic_approval_tool(make_office):
    from hq.tools.office import REQUEST_APPROVAL, definitions

    [d] = definitions([REQUEST_APPROVAL], captain="Stott")
    assert "newsletter" not in d["input_schema"]["properties"]["kind"]["enum"]
    assert "newsletter" not in d["description"]


# ---- outreach emails ----------------------------------------------------------------------------
from hq import outreach

MAIL = ("Dear Ms Rao,\n\nWe read your piece on small-cap semiconductors and thought our work might be useful "
        "to your readers. Conscious Investments is an independent research imprint run by an AI-assisted "
        "office and reviewed by its author. Each week we publish a short note that carries only price "
        "targets and ratings we have formally approved, and the office itself can be watched at "
        "https://consciousinvestments.org as the work happens.\n\nIf it is of interest, the latest note "
        "is on the site, and a reply is always welcome.\n\nWith thanks,\nEthan Stott")


@pytest.fixture
def mail_office(office_llm, monkeypatch):
    office, llm = office_llm
    cfg = {"client_relations": {"outreach": {"postal_address": "1 Main St, Springfield, IL 62701"}}}
    monkeypatch.setattr("hq.outreach.office_config", lambda: cfg)
    return office, llm


def draft(office, **kw):
    args = {"agent_id": "cr_associate", "to_name": "Maya Rao", "to_email": "maya@example.com",
            "org": "Chip Weekly", "segment": "business",
            "reason": "we thought our research may be useful to your readers",
            "subject": "Independent research you can watch being made", "body": MAIL}
    args.update(kw)
    return outreach.save_draft(office, **args)


def err(meta):
    return {p["code"] for p in meta["problems"] if p["level"] == "error"}


def test_clean_outreach_passes_and_footer_is_code_built(mail_office):
    office, _ = mail_office
    meta = draft(office)
    assert err(meta) == set() and meta["status"] == "draft"
    assert outreach.recheck(office, meta["id"]) == meta["problems"]
    meta = outreach.render(office, meta["id"])
    text = (office.outbox_dir / meta["files"]["email.txt"]).read_text()
    assert text.startswith("To: Maya Rao <maya@example.com>\nFrom: Conscious Investments <contact@consciousinvestments.org>")
    assert "1 Main St, Springfield" in text and '"unsubscribe"' in text and "not investment advice" in text.lower().replace("nothing here is investment advice", "not investment advice")
    assert "You are receiving this because we thought our research" in text


def test_outreach_gate(mail_office):
    office, _ = mail_office
    assert "address" in err(draft(office, to_email="not-an-email"))
    assert "address" in err(draft(office, to_email="contact@consciousinvestments.org"))
    assert "segment" in err(draft(office, segment="whale"))
    assert "reason" in err(draft(office, reason=""))
    assert "deceptive" in err(draft(office, subject="Re: our chat"))
    assert "spammy" in err(draft(office, subject="Act now!"))
    assert "link" in err(draft(office, body=MAIL + "\nhttps://evil.example.com/x"))
    assert "hype" in err(draft(office, body=MAIL.replace("useful", "a massive game changer")))
    assert "unapproved-target" in err(draft(office, body=MAIL + "\nBFLY has a price target of $20."))
    assert "length" in {p["code"] for p in draft(office, body="Hello. Please read our note.")["problems"]}   # a warning
    assert {p["code"] for p in draft(office, segment="retail")["problems"]} >= {"retail"}
    assert "advice" in err(draft(office, body=MAIL + "\nYou should buy RMBS now."))


def test_suppression_and_contact_limits(mail_office):
    office, _ = mail_office
    first = draft(office)
    outreach.set_status(office, first["id"], "approved", approved_at=__import__("time").time())
    assert "gap" in err(draft(office, subject="A second note"))          # too soon after the first
    outreach.suppress(office, "maya@example.com", "asked to stop")
    assert "suppressed" in err(draft(office, to_email="Maya@Example.com"))
    outreach.suppress(office, "@blocked.org", "domain asked")
    assert "suppressed" in err(draft(office, to_email="anyone@blocked.org"))
    with pytest.raises(ValueError):
        outreach.suppress(office, "nonsense", "x")


def test_outreach_needs_a_postal_address(office_llm, monkeypatch):
    office, _ = office_llm
    monkeypatch.setattr("hq.outreach.office_config", dict)
    meta = draft(office)
    assert "postal-address" in {p["code"] for p in outreach.recheck(office, meta["id"]) if p["level"] == "error"}


def test_outreach_tools_by_role():
    from hq.tools.office import tools_for

    names = lambda tier, aid, wing: {t.name for t in tools_for(tier, aid, wing)}
    harbor, wren = names("lead", "cr_lead", "client_relations"), names("associate", "cr_associate", "client_relations")
    assert {"outreach_context", "save_outreach", "read_outreach", "finalize_outreach", "suppress_contact"} <= harbor
    assert {"outreach_context", "save_outreach", "read_outreach"} <= wren
    assert not {"finalize_outreach", "suppress_contact"} & wren
    assert not {"save_outreach", "finalize_outreach"} & names("lead", "er_lead", "equity_research")


async def test_outreach_flow_and_approval(mail_office):
    office, llm = mail_office
    fields = {"to_name": "Maya Rao", "to_email": "maya@example.com", "org": "Chip Weekly", "segment": "business",
              "reason": "we thought our research may be useful to your readers",
              "subject": "Independent research you can watch being made", "body": MAIL}
    email = f"{office.ledger.today()}-chip-weekly"
    llm.script("cr_lead", tool_turn(("outreach_context", {"to_email": "maya@example.com"})),
               tool_turn(("save_outreach", fields)),
               tool_turn(("finalize_outreach", {"email": email, "note": "Warm intro."})),
               tool_turn(("save_outreach", fields)),           # frozen once filed
               text_turn("On your desk."), text_turn("Thanks."))
    office.assign("cr_lead", "Write to Maya Rao at Chip Weekly")
    await office.idle()

    ctx = json.loads(_results(llm, "cr_lead", 1)[0]["content"])
    assert ctx["contact_history"] == [] and ctx["suppressed"] is False and ctx["links_allowed"] == ["consciousinvestments.org"]
    saved = json.loads(_results(llm, "cr_lead", 2)[0]["content"])
    assert saved["errors"] == [] and saved["email"] == email
    assert "approval #1 is on Stott's desk" in _results(llm, "cr_lead", 3)[0]["content"]
    late = _results(llm, "cr_lead", 4)[0]
    assert late["is_error"] and "waiting for a decision" in late["content"]

    [card] = office.store.approvals("pending")
    assert card["kind"] == "outreach" and "maya@example.com" in card["summary"] and "Warm intro." in card["summary"]
    assert "Nothing is sent until you approve" in card["summary"]
    outreach.suppress(office, "maya@example.com", "asked to stop")        # opts out before he decides
    with pytest.raises(ValueError, match="no longer go out"):
        office.decide(card["id"], "approved")
    (outreach.root(office) / "suppressed.json").write_text("[]")
    office.decide(card["id"], "approved")
    meta = outreach.load(office, email)
    assert meta["status"] == "approved" and "contact@consciousinvestments.org" in meta["next_step"]
    assert office.outbox_view()["outreach"][0]["id"] == email
