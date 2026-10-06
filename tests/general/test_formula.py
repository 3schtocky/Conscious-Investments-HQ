from datetime import date

import pytest
from openpyxl import Workbook

from departments.quant.formula import FormulaError, WorkbookEvaluator, hardcoded_numbers


def _wb(**cells):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    other = wb.create_sheet("Calc Base")
    for ref, v in cells.items():
        sheet, _, coord = ref.partition("__")
        (other if sheet == "calc" else ws)[coord or sheet] = v
    return wb


def test_arithmetic_precedence_and_excel_negation():
    wb = _wb(A1=2, A2="=-A1^2", A3="=1+2*3^2", A4="=(1+2)*3", A5="=10/4-1", A6="=2^-1")
    ev = WorkbookEvaluator(wb)
    assert ev.cell("A2", "Inputs") == 4          # Excel: (-2)^2
    assert ev.cell("A3", "Inputs") == 19
    assert ev.cell("A4", "Inputs") == 9
    assert ev.cell("A5", "Inputs") == 1.5
    assert ev.cell("A6", "Inputs") == 0.5


def test_functions_ranges_strings_and_cross_sheet_refs():
    wb = _wb(A1=1, B1=2, C1=3, A2="=SUM(A1:C1)", A3="=MAX(A1:C1,7)", A4="=MIN(A1,B1)",
             A5='=IF(A1>=1,"yes","no")', A6="=AND(A1=1,OR(B1>5,C1=3))", A7="=ROUND(2.5,0)",
             A8="=ROUND(-2.5,0)", A9="=ABS(-3)", A10='="x"&A1',
             calc__B2="=Inputs!A2*2", calc__B3="='Calc Base'!B2+Inputs!C1")
    ev = WorkbookEvaluator(wb)
    assert ev.cell("A2", "Inputs") == 6 and ev.cell("A3", "Inputs") == 7
    assert ev.cell("A4", "Inputs") == 1 and ev.cell("A5", "Inputs") == "yes"
    assert ev.cell("A6", "Inputs") is True
    assert ev.cell("A7", "Inputs") == 3 and ev.cell("A8", "Inputs") == -3   # half away from 0
    assert ev.cell("A9", "Inputs") == 3 and ev.cell("A10", "Inputs") == "x1"
    assert ev.cell("'Calc Base'!B3") == 15


def test_if_is_lazy_and_dates_are_serials():
    wb = _wb(A1=0, A2="=IF(A1=0,0,1/A1)", A3=date(2026, 9, 30), A4="=A3-A5", A5=date(2026, 9, 29))
    ev = WorkbookEvaluator(wb)
    assert ev.cell("A2", "Inputs") == 0          # no #DIV/0! from the untaken branch
    assert ev.cell("A4", "Inputs") == 1


def test_errors_are_reported_with_the_cell():
    wb = _wb(A1="=A2", A2="=A1", B1="=1/0", C1="=FOO(1)")
    ev = WorkbookEvaluator(wb)
    with pytest.raises(FormulaError, match="circular"):
        ev.cell("A1", "Inputs")
    with pytest.raises(FormulaError, match=r"Inputs!B1: #DIV/0!"):
        ev.cell("B1", "Inputs")
    with pytest.raises(FormulaError, match="unsupported function FOO"):
        ev.cell("C1", "Inputs")


def test_hardcoded_number_scan():
    assert hardcoded_numbers("=C5*(1+C6)") == []
    assert hardcoded_numbers("=(C5+C6/2)/C7") == []
    assert hardcoded_numbers("=C5*1.08") == ["1.08"]
    assert hardcoded_numbers("=C5/365.25+Inputs!B12") == ["365.25"]
    assert hardcoded_numbers("=SUM(C5:G5)*12") == ["12"]
