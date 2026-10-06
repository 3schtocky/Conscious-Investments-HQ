"""Client Relations' slidedeck: the copy gate and the native PowerPoint, on the real Micron report."""

from __future__ import annotations

import copy
import json
import shutil
from pathlib import Path

import pytest
from erb import model as erb_model

from hq import deck, deckbuild, modelbrief
from hq.quant.workbook import load_assumptions
from hq.tools.desk import coverage_dir

MU = coverage_dir("MU")
COPY = Path(__file__).parent / "fixtures" / "deck_mu_copy.json"
pytestmark = pytest.mark.skipif(not (MU / "assumptions.yaml").exists(), reason="needs the Micron coverage pack")


def brief_reading(facts):
    names = [d["driver"] for d in facts["drivers"]]
    sc = facts["scenarios"]
    return {"view": ("The base case adopts the Street's revenue and fades margin, so the rating follows from the model. "
                     f"Our base target is ${sc['base']['price_target']:,.2f} against a price of ${facts['price']:,.2f}."),
            "drivers": [{"driver": n, "why": "It moves the base target more than most inputs."} for n in names[:4]],
            "scenarios": {"bear": "The cycle returns on schedule and margins fall toward past troughs.",
                          "base": "Revenue follows the Street and margins fade as the contracts age.",
                          "bull": "Supply stays short and the contracts hold margins higher for longer."},
            "sensitivity_note": "The target is most exposed to the discount rate and to the exit multiple applied to the final year.",
            "breaks": ["Quarterly revenue falls below the bear path for two quarters in a row.",
                       "The disclosed price floor implies a gross margin below 75.0% for most of FY'28.",
                       "Industry supply grows faster than we model once the new fabs ramp.",
                       "Customers delay payment so that receivables keep rising faster than revenue."]}


@pytest.fixture(scope="module")
def mu_facts(tmp_path_factory):
    return None


@pytest.fixture
def mu(make_office, tmp_path):
    frozen = tmp_path / "quant" / "MU" / "MU_model_v1_assumptions.yaml"
    frozen.parent.mkdir(parents=True)
    shutil.copy(MU / "assumptions.yaml", frozen)
    a, _ = load_assumptions(frozen)
    res = {k: erb_model.value_scenario(a, k) for k in ("bear", "base", "bull")}
    office, _ = make_office()
    office.store.add_model(ticker="MU", version=1, path=str(frozen.with_suffix(".xlsx")), created_by="quant_lead",
                           summary={"price_targets": {k: round(r["price_target"], 2) for k, r in res.items()},
                                    "total_returns": {k: r["total_return"] for k, r in res.items()},
                                    "rating": "Outperform", "price": float(a["price"]), "warnings": [], "as_of": str(a["as_of"]),
                                    "target_date": res["base"]["target_date"], "check": {"ok": True},
                                    "assumptions": str(frozen)})
    office.store.set_model_status("MU", 1, "approved")
    facts = modelbrief.ensure_facts(office, "MU", 1)["facts"]
    _, problems = modelbrief.save_reading(office, "MU", 1, brief_reading(facts), by="Sigma")
    assert problems == []
    return office


def errors_of(office, c):
    return [f"{p['where']}: {p['msg']}" for p in deck.errors(deck.check_copy(office, "MU", c))]


def test_the_micron_copy_passes_the_gate(mu):
    c = json.loads(COPY.read_text())
    assert errors_of(mu, c) == []


def test_the_gate_catches_what_the_words_must_not_do(mu):
    base = json.loads(COPY.read_text())

    def with_bullet(sid, i, **kw):
        c = copy.deepcopy(base)
        c["slides"][sid]["bullets"][i].update(kw)
        return c

    # a figure that is in neither the report nor the model
    bad = errors_of(mu, with_bullet("thesis", 0, text="MU grew revenue to $999.9 bn in FY'26."))
    assert any("999.9" in e and "not in the report" in e for e in bad)
    # an unapproved target, hype, advice
    assert any("target" in e.lower() for e in errors_of(mu, with_bullet("thesis", 3, text="Our price target for MU is $2,000.00.")))
    assert any(("Hype" in e or "Not our voice" in e) for e in errors_of(mu, with_bullet("thesis", 3, text="This is an unmissable, game changer of a stock.")))
    assert any("advice" in e.lower() for e in errors_of(mu, with_bullet("thesis", 3, text="You should buy MU now.")))
    # sources must exist
    assert any("no-source" in str(e) or "Name the source" in e for e in errors_of(mu, with_bullet("thesis", 0, source=[])))
    assert any("does not exist" in e for e in errors_of(mu, with_bullet("thesis", 0, source="report:99_nothing")))
    # structure
    c = copy.deepcopy(base)
    del c["slides"]["risks"]
    assert any("risks: Needs a headline" in e for e in errors_of(mu, c))
    c = copy.deepcopy(base)
    c["slides"]["thesis"]["headline"] = "word " * 20
    assert any("Headline is" in e for e in errors_of(mu, c))


def test_risks_must_weigh_as_much_as_the_upside(mu):
    c = json.loads(COPY.read_text())
    for b in c["slides"]["risks"]["bullets"]:
        b["detail"] = "Short."
        b["text"] = "A risk exists."
    assert any("risks have" in e for e in errors_of(mu, c))


def test_a_report_on_different_numbers_blocks_the_deck(mu, monkeypatch, tmp_path):
    d = tmp_path / "erb" / "coverage" / "MU"
    shutil.copytree(MU, d, ignore=shutil.ignore_patterns("filings", "*.docx", "*.pdf"))
    mj = json.loads((d / "model.json").read_text())
    mj["scenarios"]["base"]["price_target"] += 50
    (d / "model.json").write_text(json.dumps(mj))
    monkeypatch.setattr("hq.tools.desk.ERB_DIR", tmp_path / "erb")
    assert any("different targets" in e for e in errors_of(mu, json.loads(COPY.read_text())))


def test_the_deck_builds_as_native_editable_slides(mu, tmp_path):
    from pptx import Presentation

    c = json.loads(COPY.read_text())
    out = deckbuild.build(mu, "MU", c, tmp_path / "deck.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == len(deck.SLIDES) == 17
    charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    tables = sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)
    assert charts >= 5 and tables >= 4
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "$1,242.41" in text and "Outperform" in text and "approved" in text
    assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in list(prs.slides)[:14])


def test_the_builder_refuses_copy_with_errors(mu, tmp_path):
    c = json.loads(COPY.read_text())
    c["slides"]["thesis"]["bullets"][0]["text"] = "MU is a game changer worth $999.9 bn."
    with pytest.raises(ValueError, match="errors"):
        deckbuild.build(mu, "MU", c, tmp_path / "x.pptx")


def test_wren_and_harbor_get_the_deck_tools_and_they_work_end_to_end(mu):
    import asyncio
    from types import SimpleNamespace

    from hq.engine.guards import GuardBlock
    from hq.tools import deck as deck_tools
    from hq.tools.desk import desk_tools

    for tier in ("lead", "associate"):
        assert {"slidedeck_material", "save_slidedeck_copy", "build_slidedeck"} <= {t.name for t in desk_tools("client_relations", tier)}
    ctx = lambda who: SimpleNamespace(office=mu, agent=mu.agents[who], task={"id": 1})
    run = lambda coro: asyncio.run(coro)
    out = json.loads(run(deck_tools._slidedeck_material(ctx("cr_associate"), {"ticker": "MU"})))
    assert [s["id"] for s in out["slides"]] == deck.IDS and "risk" in " ".join(out["rules"]).lower()
    with pytest.raises(GuardBlock, match="No saved slidedeck copy"):
        run(deck_tools._build_slidedeck(ctx("cr_lead"), {"ticker": "MU"}))
    bad = json.loads(COPY.read_text())
    bad["slides"]["thesis"]["bullets"][0]["text"] = "MU is worth $999.9 bn."
    saved = json.loads(run(deck_tools._save_slidedeck_copy(ctx("cr_associate"), {"ticker": "MU", "copy": bad})))
    assert saved["errors"] and "Fix every error" in saved["next"]
    with pytest.raises(GuardBlock, match="errors"):
        run(deck_tools._build_slidedeck(ctx("cr_lead"), {"ticker": "MU"}))
    good = json.loads(run(deck_tools._save_slidedeck_copy(ctx("cr_associate"), {"ticker": "MU", "copy": COPY.read_text()})))
    assert good["errors"] == []
    built = json.loads(run(deck_tools._build_slidedeck(ctx("cr_lead"), {"ticker": "MU"})))
    assert built["slides"] == 17 and (mu.outbox_dir / built["built"]).is_file()
    assert json.loads(run(deck_tools._read_slidedeck_copy(ctx("cr_lead"), {"ticker": "MU"})))["clean"]


def test_the_package_ties_out_every_printed_figure_and_records_the_deck(mu, monkeypatch):
    from hq import deckpack

    monkeypatch.setattr(deckpack, "export_pdf", lambda pptx, pdf: pdf.write_bytes(b"%PDF") or True)
    _, problems = deck.save_copy(mu, "MU", json.loads(COPY.read_text()), by="Wren")
    assert deck.errors(problems) == []
    try:
        meta = deckpack.package(mu, "MU")
    except ValueError as e:
        raise AssertionError(f"{e}\n\n" + next((mu.outbox_dir / "slidedecks" / "MU" / mu.ledger.today()).glob("*Source-Map*.md")).read_text()[:6000]) from e
    stem = f"CI_Micron-Technology-Inc_MU_%s_{mu.ledger.today()}"
    assert meta["files"]["pptx"].endswith(stem % "Slidedeck" + ".pptx") and meta["files"]["pdf"].endswith(stem % "Slidedeck" + ".pdf")
    assert meta["files"]["source_map"].endswith(stem % "Source-Map" + ".md")
    assert meta["matched"] == meta["figures"] > 50 and set(meta["files"]) == {"pptx", "source_map", "pdf"}
    assert (mu.outbox_dir / meta["files"]["source_map"]).read_text().startswith("# MU slidedeck: source map")
    assert deckpack.load(mu, meta["id"])["status"] == "draft" and deckpack.recheck(mu, meta["id"]) == []


NOTE = ("## The idea\n\nMicron is the name we finished researching this week. Our base price target for MU is $1,242.41 and we rate it "
        "Outperform, per Quant model v1. The bear case matters as much: our bear price target for MU is $306.91.\n\n"
        "## What we're watching\n\nNothing else this week.\n")


async def finalize(mu, monkeypatch, *, issue=True):
    from types import SimpleNamespace

    from hq import deckpack, outbox
    from hq.tools import deck as deck_tools

    monkeypatch.setattr(deckpack, "export_pdf", lambda pptx, pdf: pdf.write_bytes(b"%PDF") or True)
    deck.save_copy(mu, "MU", json.loads(COPY.read_text()), by="Wren")
    meta = None
    if issue:
        meta = outbox.save_draft(mu, agent_id="cr_associate", title="Micron and the memory cycle", body=NOTE,
                                 x_post="Micron: the memory cycle at its peak, MU rated Outperform.", linkedin_post="Our note on MU.")
        assert meta["status"] == "draft", meta["problems"]
    ctx = SimpleNamespace(office=mu, agent=mu.agents["cr_lead"], task={"id": 1})
    args = {"ticker": "MU", **({"issue": meta["id"]} if issue else {})}
    out = await deck_tools._finalize_slidedeck(ctx, args)
    return meta, out


async def test_finalizing_puts_the_deck_and_its_newsletter_on_one_card(mu, monkeypatch):
    from conftest import text_turn

    from hq import deckpack, outbox

    mu.llm.script("cr_lead", text_turn("Noted."))
    meta, out = await finalize(mu, monkeypatch)
    card = next(c for c in mu.store.approvals("pending") if c["kind"] == "deck")
    assert "approval #" in out and card["payload"]["issue"] == meta["id"]
    assert "tied" in card["summary"] and "matching weekly note" in card["summary"]
    assert any(a.endswith(".pptx") for a in card["payload"]["attachments"]) and any(a.endswith("issue.html") for a in card["payload"]["attachments"])
    assert deckpack.load(mu, card["payload"]["deck"])["status"] == "awaiting"
    assert outbox.load(mu.outbox_dir, meta["id"])["status"] == "awaiting"
    mu.decide(card["id"], "approved")
    await mu.idle()
    assert deckpack.load(mu, card["payload"]["deck"])["status"] == "approved"
    assert outbox.load(mu.outbox_dir, meta["id"])["status"] == "approved"
    assert mu.outbox_view()["decks"][0]["status"] == "approved"


async def test_an_approval_is_refused_when_the_facts_moved_on_and_a_decline_is_always_allowed(mu, monkeypatch, tmp_path):
    from hq import deckpack

    await finalize(mu, monkeypatch, issue=False)
    card = next(c for c in mu.store.approvals("pending") if c["kind"] == "deck")
    assert "No matching newsletter" in card["summary"]
    mu.store._exec("UPDATE models SET summary=? WHERE ticker='MU'",
                   (json.dumps({**mu.store.approved_model("MU")["summary"], "rating": "Neutral"}),))
    with pytest.raises(ValueError, match="changed since this slidedeck was finalized"):
        mu.decide(card["id"], "approved")
    mu.decide(card["id"], "rejected")
    assert deckpack.load(mu, card["payload"]["deck"])["status"] == "rejected"


def test_only_the_lead_can_finalize_and_an_unclean_package_is_refused(mu, monkeypatch):
    import asyncio
    from types import SimpleNamespace

    from hq.engine.guards import GuardBlock
    from hq.tools import deck as deck_tools
    from hq.tools.desk import desk_tools

    assert "finalize_slidedeck" in {t.name for t in desk_tools("client_relations", "lead")}
    assert "finalize_slidedeck" not in {t.name for t in desk_tools("client_relations", "associate")}
    ctx = SimpleNamespace(office=mu, agent=mu.agents["cr_lead"], task={"id": 1})
    with pytest.raises(GuardBlock, match="No saved slidedeck copy"):
        asyncio.run(deck_tools._finalize_slidedeck(ctx, {"ticker": "MU"}))
    bad = json.loads(COPY.read_text())
    bad["slides"]["thesis"]["bullets"][0]["text"] = "MU is worth $999.9 bn."
    deck.save_copy(mu, "MU", bad, by="Wren")
    with pytest.raises(GuardBlock, match="still has errors"):
        asyncio.run(deck_tools._finalize_slidedeck(ctx, {"ticker": "MU"}))


async def test_harbor_and_wren_work_a_deck_through_the_engine(mu, monkeypatch):
    """The whole desk, scripted: Harbor delegates the copy, Wren's first save is refused by the gate and she fixes it,
    Harbor reads it back, builds it and puts it on the Captain's desk."""
    from conftest import text_turn, tool_turn

    from hq import deckpack

    monkeypatch.setattr(deckpack, "export_pdf", lambda pptx, pdf: pdf.write_bytes(b"%PDF") or True)
    good = json.loads(COPY.read_text())
    bad = copy.deepcopy(good)
    bad["slides"]["thesis"]["bullets"][0]["text"] = "MU grew revenue to $999.9 bn in FY'26."
    mu.llm.script("cr_lead",
                  tool_turn(("slidedeck_readiness", {"ticker": "MU"})),
                  tool_turn(("delegate", {"to": "cr_associate", "job": "Write the MU slidedeck copy to the outline and save it."})),
                  tool_turn(("read_slidedeck_copy", {"ticker": "MU"}), ("finalize_slidedeck", {"ticker": "MU"})),
                  text_turn("The MU deck is on the Captain's desk."))
    mu.llm.script("cr_associate",
                  tool_turn(("slidedeck_material", {"ticker": "MU"})),
                  tool_turn(("save_slidedeck_copy", {"ticker": "MU", "copy": bad})),
                  tool_turn(("save_slidedeck_copy", {"ticker": "MU", "copy": good})),
                  tool_turn(("submit_result", {"findings": "Saved, clean: 17 slides.", "confidence": "high"})))
    task = mu.store.create_task(assignee="cr_lead", assigned_by="captain", kind="assignment", title="Slidedeck: MU", body="Build it.")
    mu._schedule("cr_lead", task)
    await mu.idle()
    cards = [c for c in mu.store.approvals("pending") if c["kind"] == "deck"]
    assert len(cards) == 1 and "17 slides" in cards[0]["summary"]
    assert list((mu.outbox_dir / "slidedecks" / "MU" / mu.ledger.today()).glob("CI_*_MU_Source-Map_*.md"))
    assert mu.store.task(task)["status"] == "done"


def test_bullets_carry_no_closing_punctuation_and_text_never_hangs_a_word(mu):
    c = json.loads(COPY.read_text())
    c["slides"]["thesis"]["bullets"][0]["text"] += "."
    c["slides"]["risks"]["bullets"][0]["detail"] += ";"
    errs = errors_of(mu, c)
    assert any("thesis bullet 1" in e and "closing" in e for e in errs) and any("risks bullet 1" in e and "detail ends" in e for e in errs)
    assert deckbuild.clean_end("Revenue grew.") == "Revenue grew" and deckbuild.clean_end("a (b).") == "a (b)"


def test_text_never_ends_on_a_stub_line():
    from hq import deckwrap

    if deckwrap._font("Calibri", False, 17) is None:
        pytest.skip("needs the Calibri font file from PowerPoint")
    sentence = "In the year to Sep'22, the last peak, LTM EV/EBITDA had a median of 4.9x, and our 6.0x exit multiple sits inside that range"
    fixed = 0
    for tenth in range(30, 90):                      # every box width from 3.0 to 8.9 inches
        runs = deckwrap.balance([(sentence, False)], font="Calibri", size=17, width_in=tenth / 10)
        assert runs[0][0].replace(deckwrap.BREAK, " ") == sentence                     # the words never change
        lines = runs[0][0].split(deckwrap.BREAK)
        if len(lines) > 1:                           # a break was needed, and it leaves a last line of at least three words
            fixed += 1
            assert len(lines[-1].split()) >= deckwrap.MIN_LAST
    assert fixed > 0                                 # at least one width needed the fix
    narrow = deckwrap.balance([(sentence, False)], font="Calibri", size=17, width_in=6.2)
    assert len(narrow[0][0].split(deckwrap.BREAK)[-1].split()) >= deckwrap.MIN_LAST
    short = "One line of text only"
    assert deckwrap.balance([(short, False)], font="Calibri", size=17, width_in=5.0) == [(short, False)]
    # a break between a bold lead and the rest of an inline bullet
    mixed = deckwrap.balance([("Report:", True), (" ", False), ("Company facts and history are cited there to filings, with source tags", False)],
                             font="Calibri", size=15, width_in=7.0)
    assert "".join(t for t, _ in mixed).replace(deckwrap.BREAK, " ").split() == ["Report:", "Company", "facts", "and", "history", "are", "cited", "there", "to", "filings,", "with", "source", "tags"]
