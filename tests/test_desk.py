"""Phase 4 desk tools: files, the Quant workbook tools, the model registry, memory."""

from __future__ import annotations

import json

import pytest
from conftest import register_model, text_turn, tool_turn
from openpyxl import load_workbook

from HQ.tools.desk import coverage_dir

HAS_META = (coverage_dir("META") / "assumptions.yaml").exists()
needs_meta = pytest.mark.skipif(not HAS_META, reason="needs local META coverage in the submodule")


def _results(llm, who, call=1):
    return [c for c in llm.calls if c["who"] == who][call]["params"]["messages"][-1]["content"]


async def test_files_are_confined_and_write_rights_follow_the_wing(make_office):
    office, llm = make_office()
    llm.script("er_lead",
               tool_turn(("read_file", {"ticker": "META", "path": "../../pyproject.toml"}),
                         ("write_file", {"ticker": "ZZZZ", "path": "assumptions.yaml", "content": "x: 1"}),
                         ("write_file", {"ticker": "ZZZZ", "path": "notes/idea.md", "content": "hello"}),
                         ("read_file", {"ticker": "bad ticker!", "path": "x"})),
               text_turn("done"))
    office.assign("er_lead", "x")
    await office.idle()
    r = _results(llm, "er_lead")
    assert "outside this ticker's folders" in r[0]["content"]
    assert "belong to Quant" in r[1]["content"]
    assert r[2].get("is_error") is None and "Saved notes/idea.md" in r[2]["content"]
    assert "must be a US ticker" in r[3]["content"]
    assert (coverage_dir("ZZZZ") / "notes" / "idea.md").read_text() == "hello"
    import shutil
    shutil.rmtree(coverage_dir("ZZZZ"))


def test_tool_lists_by_wing():
    from HQ.tools.office import tools_for

    names = lambda tier, aid, wing: [t.name for t in tools_for(tier, aid, wing)]
    quill = names("lead", "er_lead", "equity_research")
    assert {"erb_facts", "erb_memo", "write_file", "get_model", "note_to_self"} <= set(quill)
    assert "build_model" not in quill
    delta = names("associate", "quant_associate", "quant")
    assert {"build_model", "run_simulations", "draft_assumptions", "submit_result"} <= set(delta)
    assert "request_approval" not in delta          # only the lead signs off
    vera = names("lead", "audit_lead", "audit")
    assert {"audit_log", "resolve_finding", "pause_agent", "file_incident", "read_file"} <= set(vera)
    assert "write_file" not in vera                 # Audit reads everything and writes nothing
    tally = names("associate", "audit_associate", "audit")
    assert "audit_log" in tally and "pause_agent" not in tally and "resolve_finding" not in tally


def test_web_search_only_for_research():
    from HQ.engine.llm import web_tools

    assert [t["type"] for t in web_tools({"id": "claude-sonnet-5-5"}, "equity_research")] == \
        ["web_search_20260209", "web_fetch_20260209"]
    assert web_tools({"id": "claude-haiku-4-5"}, "equity_research")[0]["type"] == "web_search_20250305"
    assert web_tools({"id": "claude-sonnet-5-5"}, "quant") == []


@needs_meta
async def test_build_model_registers_a_checked_draft_and_simulations_land_in_the_workbook(make_office):
    office, llm = make_office()
    llm.script("quant_associate", tool_turn(("build_model", {"ticker": "META"})),
               tool_turn(("run_simulations", {"ticker": "META", "runs": 300})),
               text_turn("built"))
    office.assign("quant_associate", "Build META")
    await office.idle()
    built = json.loads(_results(llm, "quant_associate", 1)[0]["content"])
    assert built["version"] == 1 and "PASSED" in built["check"]
    assert built["price_targets"]["base"] == pytest.approx(951.30, abs=0.01)   # the published model
    [m] = office.store.models("META")
    assert m["status"] == "draft" and m["summary"]["check"]["ok"]
    wb = load_workbook(m["path"])
    assert {"Outputs", "Inputs", "Calc Base", "Calc Bull", "Calc Bear", "Monte Carlo",
            "Value Drivers"} <= set(wb.sheetnames)
    sims = json.loads(_results(llm, "quant_associate", 2)[0]["content"])
    assert sims["runs"] == 300 and len(sims["top_value_drivers"]) == 4
    # a Quant workbook shows up in Delta's document history with a download link
    doc = next(d for d in office.documents("quant_associate") if d["source"] == "model")
    assert doc["file"] == "/files/quant/META/META_model_v1.xlsx"


async def test_model_approval_needs_a_passing_registered_version(make_office):
    office, llm = make_office()
    register_model(office, "RMBS", 1, ok=False)
    llm.script("quant_lead",
               tool_turn(("request_approval", {"kind": "model", "ticker": "RMBS", "version": 1,
                                               "title": "t", "summary": "s"}),
                         ("request_approval", {"kind": "model", "ticker": "RMBS", "version": 9,
                                               "title": "t", "summary": "s"})),
               text_turn("ok"))
    office.assign("quant_lead", "x")
    await office.idle()
    r = _results(llm, "quant_lead")
    assert "failed its formula check" in r[0]["content"] and "No model v9" in r[1]["content"]
    assert office.store.approvals() == []


async def test_only_approved_models_are_official(make_office):
    office, llm = make_office()
    register_model(office, "RMBS", 1)
    register_model(office, "RMBS", 2)
    llm.script("er_lead", tool_turn(("get_model", {"ticker": "RMBS"})), text_turn("ok"))
    llm.script("quant_lead", text_turn("noted"), text_turn("noted"))
    office.assign("er_lead", "check")
    await office.idle()
    assert "No approved model for RMBS" in _results(llm, "er_lead", 1)[0]["content"]
    aid = office.request_approval("quant_lead", kind="model", title="v2", summary="s",
                                  payload={"ticker": "RMBS", "version": 2}, task_id=None)
    office.decide(aid, "approved")
    llm.script("er_lead", tool_turn(("get_model", {"ticker": "RMBS"})), text_turn("ok"))
    office.assign("er_lead", "check again")
    await office.idle()
    got = json.loads(_results(llm, "er_lead", 3)[0]["content"])   # the second task's first result
    assert got["version"] == 2 and got["official"] and got["cite_as"].startswith("Quant model v2, approved")
    # an old card can't roll the official numbers back to an earlier version
    aid1 = office.request_approval("quant_lead", kind="model", title="v1", summary="s",
                                   payload={"ticker": "RMBS", "version": 1}, task_id=None)
    with pytest.raises(ValueError, match="would roll it back"):
        office.decide(aid1, "approved")
    register_model(office, "RMBS", 3)
    aid3 = office.request_approval("quant_lead", kind="model", title="v3", summary="s",
                                   payload={"ticker": "RMBS", "version": 3}, task_id=None)
    office.decide(aid3, "approved")
    statuses = {m["version"]: m["status"] for m in office.store.models("RMBS")}
    assert statuses == {1: "awaiting", 2: "superseded", 3: "approved"}


async def test_desk_notes_and_wiki_reach_the_next_task(make_office):
    office, llm = make_office()
    (office.memory_dir).mkdir(parents=True, exist_ok=True)
    (office.memory_dir / "wiki.md").write_text("- Stott prefers ranges over point estimates.")
    llm.script("quant_lead", tool_turn(("note_to_self", {"note": "Stott wants the bear case first."})),
               text_turn("noted"), text_turn("second task"))
    office.assign("quant_lead", "first")
    await office.idle()
    office.assign("quant_lead", "second")
    await office.idle()
    system = [c for c in llm.calls if c["who"] == "quant_lead"][-1]["params"]["system"]
    assert "Stott prefers ranges over point estimates." in system
    assert "Stott wants the bear case first." in system


@needs_meta
def test_file_download_endpoint_is_confined(make_office):
    from fastapi.testclient import TestClient

    from HQ.server import create_app

    office, _ = make_office()
    (office.quant_dir / "META").mkdir(parents=True)
    (office.quant_dir / "META" / "META_model_v1.xlsx").write_bytes(b"PK fake")
    with TestClient(create_app(office_factory=lambda: office)) as c:
        assert c.get("/files/quant/META/META_model_v1.xlsx").content == b"PK fake"
        assert c.get("/files/quant/META/../../../pyproject.toml").status_code == 404
        assert c.get("/files/coverage/META/assumptions.yaml").status_code == 200
        assert c.get("/files/elsewhere/META/x").status_code == 404


# Regression tests for the Phase 4 code review ---------------------------------------------
def test_half_day_horizons_match_the_engine(tmp_path):
    import yaml

    from departments.quant.workbook import build_model
    if not HAS_META:
        pytest.skip("needs META coverage")
    a = yaml.safe_load((coverage_dir("META") / "assumptions.yaml").read_text())
    a["horizon_months"] = 24   # 730.5 days: Excel ROUND and Python round() disagree
    (tmp_path / "a.yaml").write_text(yaml.safe_dump(a))
    assert build_model(tmp_path / "a.yaml", tmp_path / "m.xlsx", version=1)["ok"]


async def test_tool_failures_come_back_as_errors_not_dead_tasks(make_office):
    office, llm = make_office()
    llm.script("er_lead", tool_turn(("read_file", {"ticker": "META", "path": "brief.md", "offset": "ten"})),
               text_turn("recovered"))
    tid = office.assign("er_lead", "x")
    await office.idle()
    r = _results(llm, "er_lead")[0]
    assert r["is_error"] and "The tool failed: ValueError" in r["content"]
    assert office.store.task(tid)["status"] == "done"


async def test_research_cannot_see_quant_drafts(make_office):
    office, llm = make_office()
    register_model(office, "RMBS", 1)
    llm.script("er_lead", tool_turn(("get_model", {"ticker": "RMBS", "version": 1}),
                                  ("read_file", {"ticker": "RMBS", "path": "quant/notes/x.md"})),
               text_turn("ok"))
    office.assign("er_lead", "x")
    await office.idle()
    r = _results(llm, "er_lead")
    assert "draft" in r[0]["content"] and "private until a model is approved" in r[1]["content"]


@needs_meta
async def test_simulations_only_touch_drafts_and_builds_never_collide(make_office):
    import asyncio

    office, llm = make_office()
    llm.script("quant_associate", tool_turn(("build_model", {"ticker": "META"})), text_turn("a"))
    llm.script("quant_lead", tool_turn(("build_model", {"ticker": "META"})), text_turn("b"))
    office.assign("quant_associate", "build")
    office.assign("quant_lead", "build too")
    await office.idle()
    assert [m["version"] for m in office.store.models("META")] == [1, 2]
    aid = office.request_approval("quant_lead", kind="model", title="v2", summary="s",
                                  payload={"ticker": "META", "version": 2}, task_id=None)
    del aid
    llm.script("quant_associate", tool_turn(("run_simulations", {"ticker": "META"})), text_turn("c"))
    office.assign("quant_associate", "simulate")
    await office.idle()
    assert "frozen" in _results(llm, "quant_associate", 3)[0]["content"]
    await asyncio.sleep(0)


def test_monte_carlo_counts_growth_paths_as_uncertainty():
    from departments.quant.simulate import sigmas
    a = {"projection_years": 3, "revenue": {"base": 100, "growth": [0.10, 0.10, 0.10]},
         "scenarios": {"bull": {"growth_delta": 0.02}, "bear": {"growth_path": [-0.05, 0.0, 0.05]}}}
    assert sigmas(a)["growth_delta"] > 0.03   # the bear path is 10 pts below base on average


async def test_desk_notes_have_their_own_allowance(make_office):
    office, llm = make_office()
    llm.script("er_lead", tool_turn(("post_to_group", {"group": "stott", "text": "one"}),
                                  ("post_to_group", {"group": "stott", "text": "two different"}),
                                  ("note_to_self", {"note": "still allowed"})),
               text_turn("ok"))
    office.assign("er_lead", "x")
    await office.idle()
    assert "Saved to your desk notes" in _results(llm, "er_lead")[2]["content"]
