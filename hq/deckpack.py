"""Packaging the slidedeck: the built .pptx, Tally's source map, the PDF, and the Outbox record.

The source map ties every figure printed on every slide (code-built and written alike) to a
source in the ticker's pool (the approved model, Quant's Model Brief, the facts, filings), using
the same code Audit's final audit uses (`hq.sourcemap`). A figure that ties to nothing is an
exception and stops the package. The PDF comes from PowerPoint on this Mac; without it the
package simply has no PDF. Nothing here calls the Anthropic API.
"""

from __future__ import annotations

import json
import subprocess
import sys
import time
from pathlib import Path

from hq import deck, deckbuild, finalaudit, handoff
from hq.engine.runtime import Office
from hq.sourcemap import extract_figures, tie_out

POWERPOINT = Path("/Applications/Microsoft PowerPoint.app")
_EXPORT = '''on run argv
  set src to POSIX file (item 1 of argv)
  set dst to (item 2 of argv)
  tell application "Microsoft PowerPoint"
    open src
    set p to active presentation
    save p in (POSIX file dst) as save as PDF
    close p saving no
  end tell
end run
'''
LOCKED = {"awaiting": "waiting for a decision", "approved": "approved"}


def meta_path(office: Office, ticker: str, day: str | None = None) -> Path:
    return deck.deck_dir(office, ticker, day) / "meta.json"


def load(office: Office, deck_id: str) -> dict | None:
    """A slidedeck by id, `<TICKER>-<yyyy-mm-dd>`."""
    ticker, _, day = deck_id.partition("-")
    if not ticker or not day:
        return None
    p = meta_path(office, ticker, day)
    return json.loads(p.read_text()) if p.is_file() else None


def set_status(office: Office, deck_id: str, status: str, **extra) -> dict:
    meta = load(office, deck_id)
    if meta is None:
        raise KeyError(f"No slidedeck {deck_id!r} in the Outbox.")
    meta.update(status=status, updated=time.time(), **extra)
    meta_path(office, meta["ticker"], meta["date"]).write_text(json.dumps(meta, indent=1))
    return meta


def decks(office: Office) -> list[dict]:
    out = []
    for p in sorted((office.outbox_dir / "slidedecks").glob("*/*/meta.json")):
        try:
            out.append(json.loads(p.read_text()))
        except ValueError:
            continue
    return sorted(out, key=lambda m: m.get("updated", 0), reverse=True)


# ---- the source map ------------------------------------------------------------------------------
def slide_lines(pptx: Path) -> list[list[str]]:
    """The printed text of each slide, one line per paragraph or table row (what a reader sees)."""
    from pptx import Presentation

    out = []
    for slide in Presentation(pptx).slides:
        lines: list[str] = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                lines += [p.text.replace("\u00a0", " ").replace("\v", " ") for p in sh.text_frame.paragraphs if p.text.strip()]
            elif sh.has_table:
                lines += ["; ".join(c.text for c in row.cells) for row in sh.table.rows]
        out.append(lines)
    return out


def source_map(office: Office, ticker: str, pptx: Path) -> dict:
    pool = finalaudit.source_pool(office, ticker)
    slides, total, matched = [], 0, 0
    for n, lines in enumerate(slide_lines(pptx), start=1):
        ties = tie_out(extract_figures("\n".join(lines)), pool)
        total += len(ties)
        matched += sum(1 for t in ties if t.source)
        slides.append({"slide": n, "headline": lines[0] if lines else "",
                       "figures": [{"figure": t.figure.text, "source": t.source, "sentence": t.figure.sentence} for t in ties]})
    exceptions = [{"slide": s["slide"], "figure": f["figure"], "sentence": f["sentence"]}
                  for s in slides for f in s["figures"] if f["source"] is None]
    return {"ticker": ticker, "figures": total, "matched": matched, "slides": slides, "exceptions": exceptions}


def render_map(sm: dict, *, version: int, day: str) -> str:
    lines = [f"# {sm['ticker']} slidedeck: source map", "",
             (f"Quant model v{version}, {day}. Tally's code tied {sm['matched']} of {sm['figures']} printed figures to a source "
              f"({len(sm['exceptions'])} exceptions)."), ""]
    for s in sm["slides"]:
        if not s["figures"]:
            continue
        lines += [f"## Slide {s['slide']}", ""]
        lines += [f"- {f['figure']}: {f['source'] or 'NO SOURCE FOUND'}" for f in s["figures"]]
        lines.append("")
    return "\n".join(lines)


# ---- the PDF -------------------------------------------------------------------------------------
def export_pdf(pptx: Path, pdf: Path) -> bool:
    """PDF through PowerPoint (macOS). False when PowerPoint isn't there or the export fails."""
    if sys.platform != "darwin" or not POWERPOINT.exists():
        return False
    try:
        subprocess.run(["osascript", "-e", _EXPORT, str(pptx), str(pdf)], check=True, capture_output=True, timeout=240)
    except (subprocess.SubprocessError, OSError):
        return False
    return pdf.is_file()


# ---- the package ---------------------------------------------------------------------------------
def package(office: Office, ticker: str, *, pdf: bool = True) -> dict:
    """Build the slidedeck from today's clean copy, tie out every figure, export the PDF, record it in the Outbox."""
    saved = deck.load_copy(office, ticker)
    if saved is None:
        raise ValueError(f"No saved slidedeck copy for {ticker} today. Save one with save_slidedeck_copy and build_slidedeck first.")
    if not saved["clean"]:
        raise ValueError("Today's saved copy still has errors; fix them with save_slidedeck_copy.")
    r = handoff.readiness(office, ticker)
    if not all(c["ok"] for c in r["conditions"]):
        raise ValueError("Not ready for a slidedeck: " + "; ".join(c["detail"] for c in r["conditions"] if not c["ok"]))
    day = office.ledger.today()
    current = load(office, f"{ticker}-{day}")
    if current and current["status"] in LOCKED:
        raise ValueError(f"Today's {ticker} slidedeck is {LOCKED[current['status']]}; it can't be replaced. If something changed, "
                         "ask for the card to be sent back first.")
    d = deck.deck_dir(office, ticker, day)
    base = deck.file_stem(office, ticker, "Slidedeck", day)
    map_name = deck.file_stem(office, ticker, "Source-Map", day)
    pptx = deckbuild.build(office, ticker, saved["copy"], d / f"{base}.pptx")
    sm = source_map(office, ticker, pptx)
    version = r["version"]
    (d / f"{map_name}.json").write_text(json.dumps(sm, indent=1))
    (d / f"{map_name}.md").write_text(render_map(sm, version=version, day=day))
    if sm["exceptions"]:
        listed = "; ".join(f"slide {e['slide']} {e['figure']}" for e in sm["exceptions"][:6])
        raise ValueError(f"{len(sm['exceptions'])} printed figures tie to no source (see the source map file): {listed}. "
                         "Fix the words or ask Research and Quant to put the figure in their files.")
    files = {"pptx": f"slidedecks/{ticker}/{day}/{pptx.name}", "source_map": f"slidedecks/{ticker}/{day}/{map_name}.md"}
    if pdf and export_pdf(pptx, d / f"{base}.pdf"):
        files["pdf"] = f"slidedecks/{ticker}/{day}/{base}.pdf"
    meta = {"id": f"{ticker}-{day}", "ticker": ticker, "date": day, "title": f"{ticker} slidedeck", "status": "draft",
            "files": files, "model_version": version, "slides": len(deck.SLIDES), "figures": sm["figures"],
            "matched": sm["matched"], "by": saved["by"], "approval_id": None, "issue": None, "updated": time.time()}
    meta_path(office, ticker, day).write_text(json.dumps(meta, indent=1))
    return meta


def recheck(office: Office, deck_id: str) -> list[str]:
    """Why this slidedeck can no longer go out as it stands (empty = fine): the copy is re-gated against today's facts."""
    meta = load(office, deck_id)
    if meta is None:
        return ["the slidedeck's files are missing from the Outbox"]
    t = meta["ticker"]
    problems = []
    model = office.store.approved_model(t)
    if model is None or model["version"] != meta["model_version"]:
        problems.append(f"the approved model is no longer v{meta['model_version']}")
    elif not handoff.readiness_holds(office, t, model["version"]):
        problems.append("the rating, report or Model Brief changed since the slidedeck was finalized")
    saved = deck.load_copy(office, t, meta["date"])
    if saved is None:
        problems.append("the saved copy is missing")
    else:
        problems += [f"{p['where']}: {p['msg']}" for p in deck.errors(deck.check_copy(office, t, saved["copy"]))[:3]]
    return problems
