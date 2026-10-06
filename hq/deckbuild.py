"""Renders the slidedeck as a native, editable PowerPoint (python-pptx).

Every number on a data slide comes from the approved model's Model Brief facts; the words come
from the copy that passed `hq.deck.check_copy`. Charts are native PowerPoint charts, so an editor
can restyle them. Speaker notes carry the talking points and the sources.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from hq import deck, deckwrap, modelbrief, outbox
from hq.engine.runtime import Office
from hq.tools import desk

W, H = 13.333, 7.5
PAPER, CARD, INK, MUTED = "F7F5EE", "ECEAE1", "2E2C29", "5F5B55"
BLUE, ORANGE, GREEN, PURPLE, RED = "2A68A8", "D27030", "3A9A7A", "8A6DB8", "B5483A"
BRAND = (BLUE, ORANGE, GREEN, PURPLE)
CASE = {"bear": RED, "base": BLUE, "bull": GREEN}
HEAD, BODY = "Georgia", "Calibri"
LEFT, RIGHT = 0.6, W - 0.6
METHOD_LABEL = {"dcf": "DCF", "pe": "P/E", "ev_ebitda": "EV/EBITDA"}
PICTURES = {"business": "mix_product.png", "mispricing": "band_ltm_pe.png", "peers": "band_ltm_ev_ebitda.png"}


def add_runs(p, runs: list[tuple[str, bool]], *, size, color, font, italic=False) -> None:
    """Runs of (text, bold) into a paragraph; a BREAK inside a run becomes a line break."""
    for text, bold in runs:
        for n, chunk in enumerate(text.split(deckwrap.BREAK)):
            if n:
                p.add_line_break()
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
            r.font.color.rgb = rgb(color)


def clean_end(s: str) -> str:
    """A bullet or headline carries no closing punctuation."""
    return s.strip().rstrip(".;:,").rstrip()


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def usd(x, d=2) -> str:
    return "n/a" if x is None else f"${x:,.{d}f}"


def pct(x, d=1, sign=True) -> str:
    return "n/a" if x is None else f"{x * 100:{'+' if sign else ''}.{d}f}%"


def bn(x) -> str:
    return "n/a" if x is None else f"${x / 1e9:,.1f} bn"


# ---- drawing helpers ---------------------------------------------------------------------------
def text(slide, x, y, w, h, s, *, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, space_after=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = deckwrap.balance([(line, bold)], font=font, size=size, width_in=w)
        add_runs(p, runs, size=size, color=color, font=font, italic=italic)
    return box


def rect(slide, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _hanging_bullet(p, size: int) -> None:
    """A real bullet with a hanging indent, so wrapped lines and the detail line align under the text."""
    from lxml import etree

    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.3))))
    pPr.set("indent", str(-int(Inches(0.3))))
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for tag in ("buClr", "buFont", "buChar"):
        for old in pPr.findall(f"{{{ns}}}{tag}"):
            pPr.remove(old)
    clr = etree.SubElement(pPr, f"{{{ns}}}buClr")
    etree.SubElement(clr, f"{{{ns}}}srgbClr").set("val", ORANGE)
    etree.SubElement(pPr, f"{{{ns}}}buFont").set("typeface", "Arial")
    etree.SubElement(pPr, f"{{{ns}}}buChar").set("char", "\u2022")


def bullets(slide, x, y, w, h, items: list[tuple[str, str | None]], *, size=17, gap=10, stack=False):
    """Bulleted paragraphs, with no closing punctuation. (lead, rest) puts the lead in bold; with `stack` the
    rest sits on its own line under the lead, otherwise it follows the lead. No paragraph ends on a stub line."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    width = w - 0.3   # the hanging indent
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        _hanging_bullet(p, size)
        rest = clean_end(rest) if rest else None
        if not rest or stack:
            lead = clean_end(lead)   # a lead that introduces inline text keeps its colon
        if not rest:
            runs = [(lead, False)]
        elif stack:
            runs = [(lead, True), (deckwrap.BREAK + rest, False)]
        else:
            runs = [(lead, True), (" ", False), (rest, False)]
        if stack and rest:
            lead_runs = deckwrap.balance([(lead, True)], font=BODY, size=size, width_in=w, indent_in=0.3)
            rest_runs = deckwrap.balance([(rest, False)], font=BODY, size=size, width_in=w, indent_in=0.3)
            runs = [*lead_runs, (deckwrap.BREAK, False), *rest_runs]
        else:
            runs = deckwrap.balance(runs, font=BODY, size=size, width_in=w, indent_in=0.3)
        add_runs(p, runs, size=size, color=INK, font=BODY)
    del width
    return box


def new_slide(prs, headline: str, n: int, footer: str, *, kicker: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    if kicker:
        text(s, LEFT, 0.38, 8, 0.3, kicker.upper(), size=11, bold=True, color=MUTED)
    text(s, LEFT, 0.66, RIGHT - LEFT, 1.1, clean_end(headline), size=27, bold=True, font=HEAD, anchor=MSO_ANCHOR.TOP)
    rect(s, 0, H - 0.09, W, 0.09, BLUE)
    for i, c in enumerate(BRAND[1:], start=1):
        rect(s, W * i / 4, H - 0.09, W / 4, 0.09, c)
    text(s, LEFT, H - 0.42, 10, 0.25, footer, size=10, color=MUTED)
    text(s, RIGHT - 1, H - 0.42, 1, 0.25, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def notes(slide, body: str, sources: list[str]):
    parts = [body.strip()] if body else []
    if sources:
        parts.append("Sources: " + "; ".join(sources))
    slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)


def style_chart(chart, *, legend=False, size=11):
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = rgb(INK)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.has_title = False


def color_points(series, colors: list[str]):
    for i, c in enumerate(colors):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)


def table(slide, x, y, w, h, rows: list[list[str]], *, col_w: list[float] | None = None, size=13,
          head_fill=INK, fills: dict[tuple[int, int], str] | None = None, bold_rows: tuple[int, ...] = ()):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    t = shape.table
    tblPr = shape._element.graphic.graphicData.tbl.tblPr   # drop the default blue banding
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Emu(int(Inches(h) / len(rows)))
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = head_fill if r == 0 else (fills or {}).get((r, c), CARD if r % 2 else PAPER)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            tf = cell.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = str(val)
            run.font.size, run.font.name = Pt(size), BODY
            run.font.bold = r == 0 or r in bold_rows
            run.font.color.rgb = rgb("FFFFFF" if r == 0 else INK)
    return t


# ---- the slides ----------------------------------------------------------------------------------
class Ctx:
    def __init__(self, office: Office, ticker: str, copy: dict):
        self.office, self.ticker, self.copy = office, ticker, copy
        self.model = office.store.approved_model(ticker)
        self.brief = modelbrief.load(office, ticker, self.model["version"])
        self.f = self.brief["facts"]
        self.reading = self.brief["reading"]
        self.cdir = desk.coverage_dir(ticker)
        mj = self.cdir / "model.json"
        self.report = json.loads(mj.read_text()) if mj.is_file() else {}
        self.name = (self.report.get("name") or ticker).title().replace("Inc", "Inc.")
        approved = (datetime.fromtimestamp(self.model["approved_at"], office.ledger.tz).date().isoformat()
                    if self.model["approved_at"] else office.ledger.today())
        self.approved = approved
        self.footer = f"Conscious Investments  |  {ticker}  |  Quant model v{self.model['version']}, approved {approved}"
        self.n = 0

    def slide(self, sid: str, kicker: str = ""):
        self.n += 1
        c = (self.copy["slides"].get(sid) or {})
        self.cur = c
        return new_slide(self.prs, c.get("headline", ""), self.n, self.footer, kicker=kicker)

    def finish(self, slide, sources: list[str] | None = None):
        notes(slide, self.cur.get("notes", ""), sources or [])


def cover(c: Ctx):
    prs = c.prs
    s = prs.slides.add_slide(prs.slide_layouts[6])
    c.n += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PAPER)
    try:
        from erb.config import ASSETS_DIR

        s.shapes.add_picture(str(ASSETS_DIR / "logo_transparent.png"), Inches(LEFT), Inches(0.55), height=Inches(0.8))
    except (OSError, ImportError):   # no logo file: fall back to the name in type
        text(s, LEFT, 0.6, 6, 0.4, "CONSCIOUS INVESTMENTS", size=14, bold=True, color=MUTED)
    text(s, LEFT, 2.0, 7.6, 0.4, "INITIATING COVERAGE", size=13, bold=True, color=ORANGE)
    text(s, LEFT, 2.45, 7.8, 1.9, c.name, size=44, bold=True, font=HEAD)
    text(s, LEFT, 4.35, 7.6, 0.5, f"{c.ticker}  |  {c.f['as_of']}", size=18, color=MUTED)
    text(s, LEFT, 5.0, 7.4, 1.1, c.copy["subtitle"], size=20, italic=True, font=HEAD, color=INK)
    base = c.f["scenarios"]["base"]
    rect(s, 8.75, 1.75, 3.98, 3.9, CARD)
    colour = GREEN if c.f["rating"] == "Outperform" else BLUE
    chip = rect(s, 9.05, 2.05, 1.9, 0.5, colour, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.text_frame.text = c.f["rating"]
    r = chip.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(16), True, BODY
    r.font.color.rgb = rgb("FFFFFF")
    chip.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 9.05, 2.85, 3.4, 0.3, "BASE-CASE PRICE TARGET", size=11, bold=True, color=MUTED)
    text(s, 9.05, 3.15, 3.4, 0.9, usd(base["price_target"]), size=40, bold=True, font=HEAD)
    text(s, 9.05, 4.15, 3.5, 0.9, [f"{pct(base['price_return'])} to the target over {c.f['horizon_months']} months",
                                  f"Price {usd(c.f['price'])} on {c.f['as_of']}"], size=14, color=MUTED)
    for i, col in enumerate(BRAND):
        rect(s, W * i / 4, H - 0.09, W / 4, 0.09, col)
    notes(s, f"{c.name} ({c.ticker}). Rating {c.f['rating']}, base-case price target {usd(base['price_target'])}.", [])


def bullets_slide(c: Ctx, sid: str, kicker: str):
    s = c.slide(sid, kicker)
    items = []
    srcs: list[str] = []
    for b in c.cur["bullets"]:
        if b.get("detail"):   # bold the first sentence, then the detail
            head, _, tail = b["text"].partition(". ")
            tail = tail.strip()
            if tail and not tail.endswith((".", "!", "?")):
                tail += "."   # a sentence boundary inside the bullet stays; only the very end loses its stop
            items.append((head, " ".join(x for x in (tail, b["detail"]) if x)))
        else:
            items.append((b["text"], None))
        src = b.get("source")
        srcs += [src] if isinstance(src, str) else list(src or [])
    pic = c.cdir / "charts" / PICTURES.get(sid, "-")
    wide = RIGHT - LEFT
    if sid == "thesis":
        wide = 7.6
        rect(s, 8.75, 1.95, 3.98, 4.5, CARD)
        text(s, 9.05, 2.15, 3.4, 0.3, "THE CALL", size=11, bold=True, color=MUTED)
        base = c.f["scenarios"]["base"]
        text(s, 9.05, 2.5, 3.4, 0.5, c.f["rating"], size=26, bold=True, font=HEAD, color=GREEN)
        text(s, 9.05, 3.2, 3.4, 0.3, "BASE PRICE TARGET", size=11, bold=True, color=MUTED)
        text(s, 9.05, 3.5, 3.4, 0.7, usd(base["price_target"]), size=34, bold=True, font=HEAD)
        text(s, 9.05, 4.4, 3.5, 1.6, [f"{pct(base['price_return'])} to target", f"Bull {usd(c.f['scenarios']['bull']['price_target'])}",
                                     f"Bear {usd(c.f['scenarios']['bear']['price_target'])}", f"Price {usd(c.f['price'])}"],
             size=14, color=INK)
    elif pic.is_file():
        wide = 6.2
        s.shapes.add_picture(str(pic), Inches(7.1), Inches(1.95), width=Inches(RIGHT - 7.1))
    bullets(s, LEFT, 1.95, wide, 4.7, items, size=17 if len(items) <= 4 else 16, stack=any(b.get("detail") for b in c.cur["bullets"]))
    text(s, LEFT, H - 0.78, 9, 0.25, "Sources: " + ", ".join(sorted({x.replace("report:", "report ").replace("brief:", "brief ") for x in srcs})),
         size=10, color=MUTED, italic=True)
    c.finish(s, sorted(set(srcs)))


def targets(c: Ctx):
    s = c.slide("call", "The call")
    f = c.f
    for i, k in enumerate(modelbrief.SCENARIOS):
        x = LEFT + i * 2.55
        sc = f["scenarios"][k]
        rect(s, x, 1.95, 2.35, 2.0, CARD)
        rect(s, x, 1.95, 2.35, 0.08, CASE[k])
        text(s, x + 0.2, 2.15, 2, 0.3, k.upper() + " CASE", size=11, bold=True, color=MUTED)
        text(s, x + 0.2, 2.5, 2, 0.7, usd(sc["price_target"]), size=28, bold=True, font=HEAD)
        text(s, x + 0.2, 3.3, 2.1, 0.4, f"{pct(sc['price_return'])} to target", size=14, color=CASE[k], bold=True)
    text(s, LEFT, 4.25, 7.4, 1.6,
         [f"Price {usd(f['price'])} on {f['as_of']}. Targets are for {f['target_date']}, a {f['horizon_months']}-month horizon.",
          (f"The base case's total return, which adds the dividend, is {pct(f['scenarios']['base']['total_return'])}. The S&P 500's expected "
           f"return over the horizon is {pct(f['benchmark_return'], sign=False)}, so the base case beats it by {f['excess_return'] * 100:.1f} "
           f"points; Outperform needs more than {f['band'] * 100:.0f}.")],
         size=14, color=INK)
    if c.cur.get("note"):
        text(s, LEFT, 5.7, 7.4, 0.9, c.cur["note"], size=14, italic=True, color=MUTED)
    cd = CategoryChartData()
    cd.categories = ["Price", "Bear", "Base", "Bull"]
    cd.add_series("Price target ($)", [f["price"]] + [f["scenarios"][k]["price_target"] for k in modelbrief.SCENARIOS])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8.3), Inches(1.9), Inches(4.45), Inches(4.4), cd)
    ch = gf.chart
    style_chart(ch)
    color_points(ch.plots[0].series[0], [MUTED, RED, BLUE, GREEN])
    ch.plots[0].gap_width = 60
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format = '"$"#,##0'
    ch.plots[0].data_labels.number_format_is_linked = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    c.finish(s, ["model:facts"])


def drivers_slide(c: Ctx):
    s = c.slide("drivers", "What moves the target")
    drv = c.f["drivers"][:5]
    cd = CategoryChartData()
    cd.categories = [d["driver"] for d in reversed(drv)]
    cd.add_series("Swing in the base target ($)", [d["swing"] for d in reversed(drv)])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(LEFT), Inches(2.0), Inches(5.9), Inches(4.2), cd)
    ch = gf.chart
    style_chart(ch)
    ch.plots[0].series[0].format.fill.solid()
    ch.plots[0].series[0].format.fill.fore_color.rgb = rgb(BLUE)
    ch.plots[0].gap_width = 50
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format = '"$"#,##0'
    ch.plots[0].data_labels.number_format_is_linked = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    text(s, LEFT, 6.25, 5.9, 0.3, "Change in the base-case target when one input moves at a time. Quant's simulation.", size=10, color=MUTED, italic=True)
    items = [(d["driver"] + ":", d["why"]) for d in c.reading["drivers"]]
    bullets(s, 6.9, 2.0, RIGHT - 6.9, 4.6, items, size=15, gap=12)
    c.finish(s, ["brief:drivers"])


def forecast(c: Ctx):
    s = c.slide("forecast", "The forecast")
    f = c.f
    base = f["projections"]["base"]
    years = [f"FY{r['fy']}E" for r in base]
    cd = CategoryChartData()
    cd.categories = years
    for k in modelbrief.SCENARIOS:
        cd.add_series(k.title(), [round(v / 1e9, 1) for v in f["projections"]["revenue"][k]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT), Inches(1.95), Inches(RIGHT - LEFT), Inches(2.75), cd)
    ch = gf.chart
    style_chart(ch, legend=True)
    for ser, k in zip(ch.plots[0].series, modelbrief.SCENARIOS, strict=True):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(CASE[k])
    ch.plots[0].gap_width = 70
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb(CARD)
    ch.value_axis.tick_labels.number_format = '"$"#,##0" bn"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    rows = [["Base case", *years],
            ["Revenue", *[bn(r["revenue"]) for r in base]],
            ["Revenue growth", *[pct(r["revenue_growth"], sign=False) for r in base]],
            ["EBITDA margin", *[pct(r["ebitda_margin"], sign=False) for r in base]],
            ["EPS (GAAP)", *[usd(r["eps"]) for r in base]],
            ["Free cash flow", *[bn(r["ufcf"]) for r in base]]]
    table(s, LEFT, 4.85, RIGHT - LEFT, 1.85, rows, col_w=[2.6] + [(RIGHT - LEFT - 2.6) / len(base)] * len(base), size=12)
    c.finish(s, ["model:projections"])


def valuation(c: Ctx):
    s = c.slide("valuation", "Valuation")
    f = c.f
    names = {"dcf": "DCF", "pe": "Forward P/E", "ev_ebitda": "EV/EBITDA", "metric": "Metric multiple"}
    methods = list(f["scenarios"]["base"]["methods"])
    wts = f["scenarios"]["base"]["weights"]
    rows = [["Method", "Weight", "Bear", "Base", "Bull"]]
    for m in methods:
        rows.append([names.get(m, m), f"{wts.get(m, 0) * 100:.0f}%", *[usd(f["scenarios"][k]["methods"].get(m)) for k in modelbrief.SCENARIOS]])
    rows.append(["Blended target", "", *[usd(f["scenarios"][k]["price_target"]) for k in modelbrief.SCENARIOS]])
    ff = (c.report or {}).get("football_field") or {}
    wide = 6.3 if ff else RIGHT - LEFT
    table(s, LEFT, 2.0, wide, 0.5 * len(rows), rows, col_w=[2.1, 0.9] + [(wide - 3.0) / 3] * 3, size=13, bold_rows=(len(rows) - 1,))
    text(s, LEFT, 2.3 + 0.5 * len(rows), wide, 1.2, c.cur.get("note", ""), size=14, italic=True, color=MUTED)
    if ff:
        labels = list(ff)
        cd = CategoryChartData()
        cd.categories = list(reversed(labels))
        cd.add_series("low", [ff[k][0] for k in reversed(labels)])
        cd.add_series("Range", [ff[k][1] - ff[k][0] for k in reversed(labels)])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(7.2), Inches(1.95), Inches(RIGHT - 7.2), Inches(3.9), cd)
        ch = gf.chart
        style_chart(ch)
        ch.plots[0].series[0].format.fill.background()
        ch.plots[0].series[1].format.fill.solid()
        ch.plots[0].series[1].format.fill.fore_color.rgb = rgb(BLUE)
        ch.plots[0].gap_width = 60
        ch.value_axis.tick_labels.number_format = '"$"#,##0'
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.value_axis.has_major_gridlines = False
        text(s, 7.2, 5.9, RIGHT - 7.2, 0.5, f"Current price {usd(f['price'])}. Ranges in dollars per share.", size=10, color=MUTED, italic=True)
    c.finish(s, ["model:facts"])


def sensitivity(c: Ctx):
    s = c.slide("sensitivity", "Sensitivity")
    sen = c.f["sensitivity"]
    price = c.f["price"]
    cols = [f"{v:g}x" if "multiple" in sen["col_label"].lower() or "EV" in sen["col_label"] else pct(v, sign=False) for v in sen["cols"]]
    rows = [[f"WACC \\ {sen['col_label']}", *cols]]
    fills = {}
    for i, (w, vals) in enumerate(zip(sen["rows_wacc"], sen["values"], strict=True), start=1):
        rows.append([pct(w, sign=False), *[usd(v, 0) for v in vals]])
        for j, v in enumerate(vals, start=1):
            fills[(i, j)] = "CFE6DC" if v >= price else "F0D6C6"
    mid = len(rows) // 2
    table(s, LEFT, 2.0, 7.7, 0.55 * len(rows), rows, size=14, fills=fills, bold_rows=(mid,))
    text(s, LEFT, 2.1 + 0.55 * len(rows), 7.7, 0.6, f"Base-case DCF value per share. Green cells sit above today's price of {usd(price)}; orange cells below it.",
         size=11, color=MUTED, italic=True)
    text(s, 8.8, 2.0, RIGHT - 8.8, 3.6, c.reading["sensitivity_note"], size=17, font=HEAD, color=INK)
    c.finish(s, ["brief:sensitivity"])


def simulation(c: Ctx):
    s = c.slide("simulation", "Range of outcomes")
    mc = c.f["monte_carlo"]
    keys = ["5", "10", "25", "50", "75", "90", "95"]
    cd = CategoryChartData()
    cd.categories = [f"P{k}" for k in keys]
    cd.add_series("Simulated base-case target ($)", [mc["percentiles"][k] for k in keys])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT), Inches(1.95), Inches(7.6), Inches(4.4), cd)
    ch = gf.chart
    style_chart(ch)
    color_points(ch.plots[0].series[0], [RED, RED, ORANGE, BLUE, ORANGE, GREEN, GREEN])
    ch.plots[0].gap_width = 50
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format = '"$"#,##0'
    ch.plots[0].data_labels.number_format_is_linked = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    rect(s, 8.7, 2.0, RIGHT - 8.7, 1.8, CARD)
    text(s, 8.95, 2.15, 3.4, 0.9, f"{mc['p_above_price'] * 100:.0f}%", size=40, bold=True, font=HEAD, color=BLUE)
    text(s, 8.95, 3.05, 3.4, 0.6, f"of runs put the target above today's price of {usd(c.f['price'])}", size=13, color=MUTED)
    rect(s, 8.7, 4.0, RIGHT - 8.7, 1.8, CARD)
    text(s, 8.95, 4.15, 3.4, 0.9, f"{mc['p_beats_benchmark'] * 100:.0f}%", size=40, bold=True, font=HEAD, color=GREEN)
    text(s, 8.95, 5.05, 3.4, 0.6, "of runs beat the S&P 500 benchmark over the horizon", size=13, color=MUTED)
    text(s, 8.7, 6.0, RIGHT - 8.7, 0.6, f"{mc['runs']:,} runs of the model with growth, margin, multiple and WACC drawn at random. P50 is the median.",
         size=10, color=MUTED, italic=True)
    c.finish(s, ["model:facts"])


def wrong(c: Ctx):
    s = c.slide("wrong", "What would prove us wrong")
    items = [(f"{i}.", b) for i, b in enumerate(c.reading["breaks"], start=1)]
    bullets(s, LEFT, 2.0, 7.0, 4.6, [(b, None) for b in c.reading["breaks"]], size=17, gap=14)
    sc = c.reading["scenarios"]
    rect(s, 8.0, 2.0, RIGHT - 8.0, 4.4, CARD)
    text(s, 8.25, 2.15, 4.3, 0.3, "THE THREE CASES, IN PLAIN ENGLISH", size=11, bold=True, color=MUTED)
    y = 2.6
    for k in modelbrief.SCENARIOS:
        text(s, 8.25, y, 4.3, 0.3, k.title(), size=14, bold=True, color=CASE[k])
        text(s, 8.25, y + 0.3, 4.3, 1.0, sc[k], size=12, color=INK)
        y += 1.25
    del items
    c.finish(s, ["brief:breaks", "brief:scenarios"])


def assumptions(c: Ctx):
    s = c.slide("assumptions", "Appendix")
    f = c.f
    inp = f["inputs"]
    base = f["projections"]["base"]
    years = [f"FY{r['fy']}E" for r in base]
    rows = [["Input", *years],
            ["Revenue growth", *[pct(r["revenue_growth"], sign=False) for r in base]],
            ["EBITDA margin", *[pct(r["ebitda_margin"], sign=False) for r in base]]]
    table(s, LEFT, 2.0, RIGHT - LEFT, 1.5, rows, col_w=[2.6] + [(RIGHT - LEFT - 2.6) / len(base)] * len(base), size=13)
    w = inp.get("weights") or {}
    lines = [("Discount rate (WACC):", pct(inp["wacc"], sign=False)),
             ("Terminal value:", f"{str(inp['terminal_method']).replace('_', ' ')}" + (f" at {inp['exit_ev_ebitda']:g}x EBITDA" if inp.get("exit_ev_ebitda") else "")),
             ("Method weights:", ", ".join(f"{METHOD_LABEL.get(k, k.upper())} {v * 100:.0f}%" for k, v in w.items())),
             ("Horizon:", f"{f['horizon_months']} months, to {f['target_date']}"),
             ("Price date:", f"{f['as_of']} at {usd(f['price'])}"),
             ("Rating rule:", f"Outperform when the base total return beats the S&P 500's expected {pct(f['benchmark_return'], sign=False)} by more than {f['band'] * 100:.0f} points")]
    bullets(s, LEFT, 3.85, RIGHT - LEFT, 2.8, lines, size=15, gap=6)
    c.finish(s, ["model:facts"])


def sources_slide(c: Ctx):
    s = c.slide("sources", "Appendix")
    used: dict[str, list[str]] = {}
    for sid, sl in c.copy["slides"].items():
        for b in sl.get("bullets") or []:
            src = b.get("source")
            for x in ([src] if isinstance(src, str) else list(src or [])):
                used.setdefault(x, [])
                if sid not in used[x]:
                    used[x].append(sid)
    reports = sorted(p.name for p in c.cdir.glob("*Initiating-Coverage*") if p.suffix == ".pdf") or \
        sorted(p.name for p in c.cdir.glob("*Initiating-Coverage*"))
    lines = [("Model:", f"Quant model v{c.model['version']}, approved {c.approved}. Every price, target, return and chart on the data slides is computed from it."),
             ("Model Brief:", f"Written by {c.reading.get('by', 'Quant')} for that version; drivers, cases and breaks are quoted from it."),
             ("Report:", (reports[0] if reports else "Initiating-coverage report") + ". Company facts and history are cited there to filings, with source tags.")]
    lines += [("Words on these slides:", ", ".join(sorted(used)) or "none")]
    bullets(s, LEFT, 2.0, RIGHT - LEFT, 4.6, lines, size=15, gap=10)
    c.finish(s, sorted(used))


def disclosures(c: Ctx):
    s = c.slide("disclosures", "Important")
    st = outbox.settings()
    text(s, LEFT, 1.95, RIGHT - LEFT, 4.6,
         [st["disclaimer"].replace("This note", "This presentation"),
          ("Ratings are relative to the author's expected return for the S&P 500 over the horizon. Outperform: expected total return exceeds the "
           "S&P 500 by more than 5 percentage points. Neutral: within 5 points. Underperform: trails by more than 5 points."),
          "Holdings disclosure: " + st["holdings_disclosure"]], size=13, color=INK, space_after=12)
    c.finish(s, [])


BUILDERS = {"cover": cover, "call": targets, "drivers": drivers_slide, "forecast": forecast, "valuation": valuation,
            "sensitivity": sensitivity, "simulation": simulation, "wrong": wrong, "assumptions": assumptions,
            "sources": sources_slide, "disclosures": disclosures}
KICKERS = {"thesis": "The case", "business": "The company", "mispricing": "The mispricing", "peers": "Peers",
           "catalysts": "Catalysts", "risks": "Risks"}


def build(office: Office, ticker: str, copy: dict, out: Path) -> Path:
    """Render the slidedeck. The caller has already run `deck.check_copy`; this refuses copy with errors."""
    problems = deck.errors(deck.check_copy(office, ticker, copy))
    if problems:
        raise ValueError("The copy has errors: " + "; ".join(f"{p['where']}: {p['msg']}" for p in problems[:5]))
    c = Ctx(office, ticker, copy)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    c.prs = prs
    for sid, _, who, _ in deck.SLIDES:
        if who == "words":
            bullets_slide(c, sid, KICKERS[sid])
        else:
            BUILDERS[sid](c)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.core_properties.title = f"{c.name} ({ticker}) initiating coverage"
    prs.core_properties.author = "Conscious Investments"
    prs.save(out)
    return out
