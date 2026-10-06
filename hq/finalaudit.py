"""The final audit: once a name's deliverables are finished, every figure is checked before the
Captain relies on them, and failures go back to their owners as targeted redos.

Plain code except for the rulings:
1. `documents` gathers the ticker's finished deliverables (research memo and brief, Quant's
   Model Brief, the initiation report, newsletter issues that name it).
2. `source_pool` gathers what figures may be traced to (the approved model, the Model Brief, the
   facts folder, filings, notes).
3. `run` ties every figure out (`hq.sourcemap`). Matches clear for free. Each sentence that holds
   an unmatched figure becomes one flag, which Vera rules on (Tally traces).
4. `request_redo` (Vera's tool) sends the owner only the failed parts, in a fresh task, and counts
   rounds per part. A part that still fails after `MAX_REDOS` goes to the Captain as an incident.
5. When a redo finishes, `after_redo` re-checks and closes what is now fixed. Nothing is rewritten
   by Audit and nothing here calls the Anthropic API.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

from hq import modelbrief
from hq.engine.runtime import OFFICE, Office
from hq.sourcemap import SourcePool, extract_figures, tie_out
from hq.tools import desk

log = logging.getLogger(__name__)
KIND = "final_audit"
RULE = "unsourced_figure"
MAX_REDOS = 2
REDO_PREFIX = "Audit redo:"
OWNERS = {"memo.md": "er_lead", "brief.md": "er_lead", "pitch.md": "screen_lead",
          "model brief": "quant_lead", "report": "er_lead", "newsletter": "cr_lead"}
SOURCE_SUFFIXES = {".csv", ".json", ".md", ".txt"}
MAX_SOURCE_BYTES = 6_000_000


@dataclass
class Doc:
    name: str       # memo.md, model brief, report, newsletter 2026-10-05-slug
    owner: str      # role id
    text: str


def _read(path: Path) -> str:
    try:
        return path.read_text(errors="replace")
    except OSError:
        return ""


def _docx_text(path: Path) -> str:
    try:
        import docx
        d = docx.Document(str(path))
    except (ImportError, ValueError, OSError, KeyError):
        return ""
    lines = [("## " if p.style.name.lower().startswith("heading") else "") + p.text
             for p in d.paragraphs if p.text.strip()]
    for t in d.tables:   # table cells are figures too; each row is a line
        lines += [" | ".join(c.text for c in row.cells) + "." for row in t.rows]
    return "\n\n".join(lines)


def documents(office: Office, ticker: str) -> list[Doc]:
    out: list[Doc] = []
    cdir = desk.coverage_dir(ticker)
    for name in ("memo.md", "brief.md", "pitch.md"):
        if (cdir / name).is_file():
            out.append(Doc(name, OWNERS[name], _read(cdir / name)))
    model = office.store.approved_model(ticker)
    if model:
        b = modelbrief.load(office, ticker, model["version"])
        if b:
            out.append(Doc("model brief", OWNERS["model brief"], modelbrief.render(b)))
    if cdir.is_dir():
        for p in sorted(cdir.glob("*Initiating-Coverage*.docx")):
            out.append(Doc("report", OWNERS["report"], _docx_text(p)))
    news = Path(office.outbox_dir) / "newsletters"
    if news.is_dir():
        for d in sorted(news.iterdir()):
            issue = d / "issue.md"
            if issue.is_file() and re.search(rf"\b{re.escape(ticker)}\b", _read(issue)):
                out.append(Doc(f"newsletter {d.name}", OWNERS["newsletter"], _read(issue)))
    return out


def source_pool(office: Office, ticker: str) -> SourcePool:
    pool = SourcePool()
    model = office.store.approved_model(ticker)
    if model:
        pool.add_json(model["summary"], "approved model")
        b = modelbrief.load(office, ticker, model["version"])
        if b:
            pool.add_json(b["facts"], "model brief")
    cdir = desk.coverage_dir(ticker)
    mj = cdir / "model.json"
    if mj.is_file() and mj.stat().st_size < MAX_SOURCE_BYTES:
        pool.add_source("model.json", _read(mj))
    for folder in ("facts", "notes"):
        base = cdir / folder
        if not base.is_dir():
            continue
        for p in sorted(base.rglob("*")):
            if p.is_file() and p.suffix in SOURCE_SUFFIXES and p.stat().st_size < MAX_SOURCE_BYTES:
                pool.add_source(f"{folder}/{p.name}", _read(p))
    for name in ("sources.md", "assumptions.yaml"):
        if (cdir / name).is_file():
            pool.add_source(name, _read(cdir / name))
    return pool


def audit_documents(office: Office, ticker: str) -> dict:
    """Tie out every document. Returns counts and the exceptions grouped by sentence."""
    pool = source_pool(office, ticker)
    docs = documents(office, ticker)
    total = matched = 0
    groups: list[dict] = []
    for d in docs:
        ties = tie_out(extract_figures(d.text), pool)
        total += len(ties)
        matched += sum(1 for t in ties if t.source)
        by_sentence: dict[str, dict] = {}
        for t in ties:
            if t.source:
                continue
            g = by_sentence.setdefault(t.figure.sentence, {
                "doc": d.name, "owner": d.owner, "where": t.figure.where,
                "sentence": t.figure.sentence, "figures": []})
            g["figures"].append(t.figure)
        groups += by_sentence.values()
    return {"ticker": ticker, "documents": [d.name for d in docs], "figures": total,
            "matched": matched, "exceptions": groups}


def _subject(ticker: str, g: dict) -> str:
    return f"{ticker} {g['doc']}" + (f" · {g['where']}" if g["where"] else "")


def _detail(g: dict) -> str:
    sent = g["sentence"] if len(g["sentence"]) <= 220 else g["sentence"][:219] + "…"
    figs = ", ".join(dict.fromkeys(f.text for f in g["figures"]))
    return f'"{sent}" has no source for: {figs}.'


def run(office: Office, ticker: str) -> dict:
    """The final audit: tie out, record a flag per failed sentence, call Vera in on them."""
    result = audit_documents(office, ticker)
    from hq.audit import Finding

    found = [Finding(RULE, "flag", _subject(ticker, g), _detail(g)) for g in result["exceptions"]]
    owner_of = {_subject(ticker, g): g["owner"] for g in result["exceptions"]}
    ids: list[int] = []
    for f in found:   # record under the owning colleague so the Audit tab says whose work it is
        ids += office.record_findings(owner_of[f.subject], None, [f], with_ids=True)
    result["findings"] = ids
    office.bus.publish("final_audit", "audit_associate" if "audit_associate" in office.agents else None,
                       None, ticker=ticker, figures=result["figures"], matched=result["matched"],
                       exceptions=len(found))
    return result


# redo loop ----------------------------------------------------------------------------------
def part_of(finding: dict) -> str:
    """The stable identity of a failed part: section plus figures, however the sentence is reworded."""
    return finding["subject"] + "|" + finding["detail"].rsplit("for: ", 1)[-1]


def redo_count(office: Office, part: str) -> int:
    return office.store.redo_count(part)


def request_redo(office: Office, finding_ids: list[int], *, by: str) -> dict:
    """Send upheld findings back to their owners: one fresh task per owner listing only the failed
    parts. A part past MAX_REDOS is not sent again; it is filed to the Captain instead."""
    sends: dict[str, list[dict]] = {}
    escalated: list[dict] = []
    for fid in finding_ids:
        f = office.store.finding(fid)
        if f["rule"] != RULE:
            raise ValueError(f"finding #{fid} is not a final-audit figure finding.")
        if f["status"] != "upheld":
            raise ValueError(f"finding #{fid} is {f['status']}: rule on it (upheld) before asking for a redo.")
        part = part_of(f)
        (escalated if redo_count(office, part) >= MAX_REDOS else sends.setdefault(f["agent"], [])).append(
            {**f, "part": part})
    tasks: list[int] = []
    for owner, items in sends.items():
        if owner not in office.agents:
            continue
        lines = "\n".join(f"- [finding #{i['id']}] {i['subject']}: {i['detail']}"
                          + (f" Correct source: {i['note']}" if i.get("note") else "") for i in items)
        body = (f"Audit upheld {len(items)} part{'s' if len(items) != 1 else ''} of your work. Fix "
                "only these, in the files they name; leave everything else as it is.\n\n"
                f"{lines}\n\nFor each: find the figure's real source (get_model, the facts folder, a "
                "filing), then correct the number or cut the claim. Do not invent a source. If the "
                "figure cannot be sourced, remove it or reword without it. Reply with one line per "
                "finding saying what you changed.")
        title = f"{REDO_PREFIX} {len(items)} part{'s' if len(items) != 1 else ''}"[:70]
        task_id = office.store.create_task(assignee=owner, assigned_by="audit_lead" if by == "audit_lead" else by,
                                           kind="assignment", title=title, body=body)
        for i in items:
            office.store.add_redo(i["part"], i["id"], task_id)
        office._schedule(owner, task_id)
        tasks.append(task_id)
    for i in escalated:
        office.raise_incident(i["agent"], None, "audit: sourcing",
                              f"Still failing after {MAX_REDOS} redos: {i['subject']}. {i['detail']} "
                              "Held for your decision.")
    return {"tasks": tasks, "escalated": [i["id"] for i in escalated]}


def after_redo(office: Office, task: dict) -> list[int]:
    """A redo finished: re-check the ticker's documents and close findings that now pass. Returns
    the ids closed. Findings that still fail stay upheld; the next request_redo counts the round."""
    if not task["title"].startswith(REDO_PREFIX):
        return []
    closed: list[int] = []
    tickers = {f["subject"].split(" ", 1)[0] for f in office.store.findings(statuses=("upheld",))
               if f["rule"] == RULE and f["agent"] == task["assignee"]}
    for t in tickers:
        still = {_subject(t, g) + "|" + ", ".join(dict.fromkeys(x.text for x in g["figures"]))
                 for g in audit_documents(office, t)["exceptions"]}
        for f in office.store.findings(statuses=("upheld",)):
            if f["rule"] != RULE or not f["subject"].startswith(f"{t} "):
                continue
            key = f["subject"] + "|" + f["detail"].rsplit("for: ", 1)[-1].rstrip(".")
            if key not in still:
                office.store.set_finding(f["id"], "cleared", by=OFFICE, note="Fixed in the redo; re-checked by code.")
                closed.append(f["id"])
    return closed


# the offer ----------------------------------------------------------------------------------
def readiness(office: Office, ticker: str) -> dict:
    model = office.store.approved_model(ticker)
    docs = documents(office, ticker) if model else []
    names = {d.name for d in docs}
    ok = model is not None and "memo.md" in names
    done = any(c for c in office.store.approvals() if c["kind"] == KIND
               and c["payload"].get("ticker") == ticker and c["payload"].get("version") == (model or {}).get("version"))
    return {"ticker": ticker, "version": model["version"] if model else None, "documents": sorted(names),
            "ready": ok and not done}


def maybe_offer(office: Office, ticker: str) -> int | None:
    """File the 'Ready for final audit' card for a name whose deliverables are done. Never raises."""
    try:
        r = readiness(office, ticker)
        if not r["ready"]:
            return None
        return office.request_approval(
            OFFICE, kind=KIND, title=f"Ready for final audit: {ticker}",
            summary=(f"{ticker}'s deliverables are finished: {', '.join(r['documents'])}. Approve to run "
                     "the final audit: code ties every figure to the approved model, the facts and the "
                     "filings, and only unmatched figures go to Audit. Anything wrong goes back to its "
                     "owner as a targeted redo. Decline to leave it."),
            task_id=None, payload={"ticker": ticker, "version": r["version"], "attachments": []})
    except Exception:
        log.exception("final-audit offer failed for %s", ticker)
        return None


def decided(office: Office, card: dict, decision: str) -> dict | None:
    if card["kind"] != KIND or decision != "approved":
        return None
    return run(office, card["payload"]["ticker"])


def offer_for_card(office: Office, card: dict) -> list[int]:
    """When the Captain approves a Client Relations piece, offer the final audit for each name it
    covers that has an approved model. A memo package names its ticker; a newsletter is read."""
    if card["kind"] not in ("newsletter", "deliverable"):
        return []
    p = card["payload"]
    names = {p["ticker"]} if p.get("ticker") else set()
    if p.get("issue"):
        issue = Path(office.outbox_dir) / "newsletters" / str(p["issue"]) / "issue.md"
        if issue.is_file():
            names |= {m["ticker"] for m in office.store.models() if m["status"] == "approved"
                      and re.search(rf"\b{re.escape(m['ticker'])}\b", _read(issue))}
    return [c for t in sorted(names) if (c := maybe_offer(office, t)) is not None]


def view(office: Office) -> dict:
    """What the Audit tab shows: each approved name, its last run and where its findings stand."""
    runs: dict[str, dict] = {}
    for e in office.store.events_where(types=["final_audit"], limit=200):
        runs[e["payload"]["ticker"]] = {"ts": e["ts"], "figures": e["payload"]["figures"],
                                        "matched": e["payload"]["matched"],
                                        "exceptions": e["payload"]["exceptions"]}
    fnd = [f for f in office.store.findings(limit=400) if f["rule"] == RULE]
    out = []
    for t in sorted({m["ticker"] for m in office.store.models() if m["status"] == "approved"}):
        cdir = desk.coverage_dir(t)
        if not (cdir / "memo.md").is_file():
            continue
        mine = [f for f in fnd if f["subject"].startswith(f"{t} ")]
        parts = [{"id": f["id"], "status": f["status"], "redos": redo_count(office, part_of(f)),
                  "held": f["status"] == "upheld" and redo_count(office, part_of(f)) >= MAX_REDOS}
                 for f in mine]
        out.append({"ticker": t, "last_run": runs.get(t), "parts": parts,
                    "open": sum(1 for p in parts if p["status"] in ("open", "reviewing", "upheld")),
                    "held": sum(1 for p in parts if p["held"])})
    return {"tickers": out, "max_redos": MAX_REDOS}


def start(office: Office, ticker: str) -> dict:
    """The Captain's Run button: only for a name with an approved model and a memo."""
    t = ticker.upper()
    if not desk.TICKER.match(t):
        raise ValueError("That is not a ticker.")
    if office.store.approved_model(t) is None or not (desk.coverage_dir(t) / "memo.md").is_file():
        raise ValueError(f"{t} needs an approved model and a finished memo before it can be audited.")
    r = run(office, t)
    return {k: r[k] for k in ("ticker", "documents", "figures", "matched", "findings")} | {
        "exceptions": len(r["exceptions"])}
