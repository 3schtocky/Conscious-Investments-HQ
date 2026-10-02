"""A small Excel formula evaluator for the workbooks Quant builds.

It exists so the office can prove a model workbook is right without Excel: every formula cell
is recomputed here and compared with the Python valuation engine (`erb.model`). It supports
exactly the subset the workbook builder emits:

    numbers, strings, TRUE/FALSE, cell refs (A1, $A$1, Sheet!A1, 'Sheet name'!A1), ranges inside
    functions (A1:E1), + - * / ^, unary -, comparisons (= <> < > <= >=), & (concatenation),
    SUM MIN MAX IF AND OR ROUND INT ABS

Excel precedence is followed, including -2^2 = 4 (negation binds tighter than ^).
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter

_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<num>\d+(?:\.\d*)?(?:[eE][-+]?\d+)?|\.\d+(?:[eE][-+]?\d+)?)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:(?:'[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<func>[A-Z][A-Z0-9.]*(?=\())
  | (?P<bool>TRUE|FALSE)
  | (?P<op><>|<=|>=|[-+*/^&=<>(),])
""", re.VERBOSE)

EXCEL_EPOCH = date(1899, 12, 30)


class FormulaError(ValueError):
    pass


def to_excel_number(v):
    """Cell values as Excel sees them: dates become serial numbers."""
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return float((v - EXCEL_EPOCH).days)
    return v


def tokenize(formula: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    while pos < len(formula):
        m = _TOKEN.match(formula, pos)
        if not m:
            raise FormulaError(f"can't read {formula[pos:pos + 20]!r} in ={formula}")
        pos = m.end()
        kind = m.lastgroup
        if kind != "ws":
            out.append((kind, m.group()))
    return out


@dataclass
class Ref:
    sheet: str | None
    col: int
    row: int
    col2: int | None = None
    row2: int | None = None


def parse_ref(text: str) -> Ref:
    sheet = None
    if "!" in text:
        sheet, text = text.rsplit("!", 1)
        sheet = sheet.strip("'")
    parts = text.replace("$", "").split(":")
    m1 = re.fullmatch(r"([A-Z]+)(\d+)", parts[0])
    ref = Ref(sheet, column_index_from_string(m1.group(1)), int(m1.group(2)))
    if len(parts) == 2:
        m2 = re.fullmatch(r"([A-Z]+)(\d+)", parts[1])
        ref.col2, ref.row2 = column_index_from_string(m2.group(1)), int(m2.group(2))
    return ref


class _Parser:
    """Recursive descent: comparison < concat < additive < multiplicative < power < unary."""

    def __init__(self, tokens, resolve):
        self.t, self.i, self.resolve = tokens, 0, resolve

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self, value=None):
        tok = self.peek()
        if value is not None and tok[1] != value:
            raise FormulaError(f"expected {value!r}, got {tok[1]!r}")
        self.i += 1
        return tok

    def parse(self):
        v = self.comparison()
        if self.i != len(self.t):
            raise FormulaError(f"unexpected {self.peek()[1]!r}")
        return v

    def comparison(self):
        v = self.concat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            r = self.concat()
            v = {"=": v == r, "<>": v != r, "<": v < r, ">": v > r, "<=": v <= r, ">=": v >= r}[op]
        return v

    def concat(self):
        v = self.additive()
        while self.peek()[1] == "&":
            self.take()
            v = f"{v}{self.additive()}"
        return v

    def additive(self):
        v = self.multiplicative()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            r = self.multiplicative()
            v = v + r if op == "+" else v - r
        return v

    def multiplicative(self):
        v = self.power()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            r = self.power()
            if op == "/" and r == 0:
                raise FormulaError("#DIV/0!")
            v = v * r if op == "*" else v / r
        return v

    def power(self):
        v = self.unary()
        while self.peek()[1] == "^":
            self.take()
            v = v ** self.unary()
        return v

    def unary(self):
        if self.peek()[1] == "-":
            self.take()
            return -self.unary()
        if self.peek()[1] == "+":
            self.take()
            return self.unary()
        return self.primary()

    def primary(self):
        kind, text = self.take()
        if kind == "num":
            return float(text)
        if kind == "str":
            return text[1:-1].replace('""', '"')
        if kind == "bool":
            return text == "TRUE"
        if kind == "ref":
            ref = parse_ref(text)
            if ref.col2 is not None:
                raise FormulaError("a range can only be used inside a function")
            return self.resolve(ref)
        if kind == "func":
            return self.call(text)
        if text == "(":
            v = self.comparison()
            self.take(")")
            return v
        raise FormulaError(f"unexpected {text!r}")

    def args(self):
        self.take("(")
        out = []
        if self.peek()[1] == ")":
            self.take()
            return out
        while True:
            kind, text = self.peek()
            if kind == "ref" and ":" in text:
                self.take()
                out.append(self.range_values(parse_ref(text)))
            else:
                out.append(self.comparison())
            if self.peek()[1] == ",":
                self.take()
                continue
            self.take(")")
            return out

    def range_values(self, ref: Ref) -> list:
        vals = []
        for r in range(ref.row, ref.row2 + 1):
            for c in range(ref.col, ref.col2 + 1):
                vals.append(self.resolve(Ref(ref.sheet, c, r)))
        return vals

    def call(self, name):
        if name == "IF":   # lazy: only the chosen branch is evaluated (like Excel)
            self.take("(")
            cond = self.comparison()
            self.take(",")
            then_at = self.i
            self._skip_arg()
            else_at = None
            if self.peek()[1] == ",":
                self.take()
                else_at = self.i
                self._skip_arg()
            self.take(")")
            after = self.i
            chosen = then_at if cond else else_at
            if chosen is None:   # IF(cond, x) with a false cond
                return False
            self.i = chosen
            v = self.comparison()
            self.i = after
            return v
        args = self.args()
        flat = [x for a in args for x in (a if isinstance(a, list) else [a])]
        nums = [x for x in flat if isinstance(x, (int, float)) and not isinstance(x, bool)]
        if name == "SUM":
            return float(sum(nums))
        if name == "MIN":
            return float(min(nums))
        if name == "MAX":
            return float(max(nums))
        if name == "AND":
            return all(bool(x) for x in flat)
        if name == "OR":
            return any(bool(x) for x in flat)
        if name == "INT":
            return float(math.floor(args[0]))
        if name == "ABS":
            return abs(args[0])
        if name == "ROUND":
            x, n = args[0], int(args[1])
            q = 10 ** n   # Excel rounds half away from zero
            return math.floor(abs(x) * q + 0.5) / q * (1 if x >= 0 else -1)
        raise FormulaError(f"unsupported function {name}")

    def _skip_arg(self) -> None:
        """Move past one argument without evaluating it."""
        depth = 0
        while True:
            text = self.peek()[1]
            if text is None:
                raise FormulaError("unterminated IF")
            if text == "(":
                depth += 1
            elif text == ")":
                if depth == 0:
                    return
                depth -= 1
            elif text == "," and depth == 0:
                return
            self.i += 1


class WorkbookEvaluator:
    """Evaluate every formula in an openpyxl workbook (loaded with data_only=False)."""

    def __init__(self, wb):
        self.wb = wb
        self.cache: dict[tuple[str, int, int], object] = {}
        self._stack: set[tuple[str, int, int]] = set()

    def value(self, sheet: str, col: int, row: int):
        key = (sheet, col, row)
        if key in self.cache:
            return self.cache[key]
        if key in self._stack:
            raise FormulaError(f"circular reference at {sheet}!{get_column_letter(col)}{row}")
        cell = self.wb[sheet].cell(row=row, column=col)
        raw = cell.value
        if isinstance(raw, str) and raw.startswith("="):
            self._stack.add(key)
            try:
                v = _Parser(tokenize(raw[1:]),
                            lambda ref: self.value(ref.sheet or sheet, ref.col, ref.row)).parse()
            except FormulaError as e:
                raise FormulaError(f"{sheet}!{cell.coordinate}: {e}") from None
            finally:
                self._stack.discard(key)
        else:
            v = to_excel_number(raw) if raw is not None else 0.0
        self.cache[key] = v
        return v

    def cell(self, ref: str, sheet: str | None = None):
        r = parse_ref(ref)
        return self.value(r.sheet or sheet, r.col, r.row)

    def evaluate_all(self) -> int:
        n = 0
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        self.value(ws.title, c.column, c.row)
                        n += 1
        return n


_LITERAL = re.compile(r"(?<![A-Z$\d.])(\d+(?:\.\d+)?)(?![\d.]*[A-Z(])")


def hardcoded_numbers(formula: str, allowed=(0.0, 1.0, 2.0)) -> list[str]:
    """Numeric literals in a formula, other than a few structural ones (0, 1, 2).

    Quant's rule: no numbers buried in calculations. Everything else must live on the
    Inputs tab and be referenced."""
    out = []
    for kind, text in tokenize(formula.lstrip("=")):
        if kind == "num" and float(text) not in allowed:
            out.append(text)
    return out
