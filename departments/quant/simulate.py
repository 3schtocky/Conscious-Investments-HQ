"""Quant's simulations: Monte Carlo on the base case and one-at-a-time value drivers.

Both run erb's own valuation engine (`erb.model.value_scenario`), so every simulated price
target is computed exactly as the model computes it. Results are written to the workbook as
clearly labelled script output (values, not formulas).

Monte Carlo: growth, margin, multiple and WACC deltas are drawn independently from normal
distributions centred on the base case, scaled so the model's own bull and bear deltas sit at
roughly the 90th and 10th percentiles (1.2816 standard deviations).
"""

from __future__ import annotations

import copy
import math

import numpy as np
from erb import model as erb_model
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

Z90 = 1.2816   # standard normal 90th percentile
KEYS = ("growth_delta", "margin_delta", "multiple_delta", "wacc_delta")
BOLD = Font(name="Calibri", bold=True)
TITLE = Font(name="Calibri", bold=True, size=14)
NOTE = Font(name="Calibri", italic=True, color="5F5B55")
HEAD_FILL = PatternFill("solid", fgColor="2E2C29")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF")


def _pt(a: dict, deltas: dict) -> float:
    a2 = copy.deepcopy(a)
    a2.setdefault("scenarios", {})["_sim"] = deltas
    return float(erb_model.value_scenario(a2, "_sim")["price_target"])


def _effective_growth_delta(a: dict, name: str) -> float:
    """A scenario's average growth shift vs base, including a custom growth path."""
    sc = erb_model.scenario(a, name)
    if sc.growth_path is None:
        return float(sc.growth_delta)
    n = int(a["projection_years"])
    segs = a["revenue"].get("segments") or {"Total": a["revenue"]}
    base = np.mean([np.mean(erb_model._vec(c["growth"], n)) for c in segs.values()])
    return float(np.mean(erb_model._vec(sc.growth_path, n)) - base + sc.growth_delta)


def sigmas(a: dict) -> dict[str, float]:
    bull = vars(erb_model.scenario(a, "bull"))
    bear = vars(erb_model.scenario(a, "bear"))
    out = {k: abs(float(bull[k]) - float(bear[k])) / (2 * Z90) for k in KEYS}
    out["growth_delta"] = abs(_effective_growth_delta(a, "bull")
                              - _effective_growth_delta(a, "bear")) / (2 * Z90)
    return out


def monte_carlo(a: dict, n: int = 2000, seed: int = 7) -> dict:
    rng = np.random.default_rng(seed)
    sd = sigmas(a)
    draws = {k: rng.normal(0.0, sd[k], n) for k in KEYS}
    price = float(a["price"])
    pts = []
    for i in range(n):
        deltas = {k: float(draws[k][i]) for k in KEYS}
        try:
            v = _pt(a, deltas)
        except (ZeroDivisionError, FloatingPointError, ValueError):
            continue
        if math.isfinite(v):
            pts.append(v)
    if len(pts) < max(20, n // 10):
        raise ValueError(f"only {len(pts)} of {n} simulations produced a valid price target; "
                         "check the assumptions (e.g. WACC close to terminal growth)")
    arr = np.array(pts)
    h = float(a.get("horizon_months", 15)) / 12
    bench = (1 + float((a.get("benchmark") or {}).get("sp500_expected_return", 0.08))) ** h - 1
    dps = float(a.get("dividend_per_share", 0.0))
    total = arr / price - 1 + dps * h / price
    pct = {p: float(np.percentile(arr, p)) for p in (5, 10, 25, 50, 75, 90, 95)}
    counts, edges = np.histogram(arr, bins=20)
    return {
        "n": len(arr), "seed": seed, "sigmas": sd, "price": price,
        "mean": float(arr.mean()), "percentiles": pct,
        "p_above_price": float((arr > price).mean()),
        "p_beats_benchmark": float((total > bench).mean()),
        "histogram": {"edges": [float(e) for e in edges], "counts": [int(c) for c in counts]},
    }


def drivers(a: dict) -> list[dict]:
    """Base price target when one input moves at a time, largest swing first."""
    base = _pt(a, {})
    out = []

    def shock(label, lo_desc, hi_desc, lo_a, hi_a):
        lo, hi = lo_a(), hi_a()
        out.append({"driver": label, "low_case": lo_desc, "high_case": hi_desc,
                    "pt_low": lo, "pt_high": hi, "swing": abs(hi - lo)})

    def with_margins(key, bump):
        def run():
            a2 = copy.deepcopy(a)
            n = int(a2["projection_years"])
            a2["margins"][key] = [v + bump for v in erb_model._vec(a2["margins"][key], n)]
            return _pt(a2, {})
        return run

    def with_wacc(bump):
        return lambda: _pt(a, {"wacc_delta": bump})

    shock("Revenue growth (every year)", "-1 pt", "+1 pt",
          lambda: _pt(a, {"growth_delta": -0.01}), lambda: _pt(a, {"growth_delta": 0.01}))
    shock("EBITDA margin (by final year)", "-1 pt", "+1 pt",
          lambda: _pt(a, {"margin_delta": -0.01}), lambda: _pt(a, {"margin_delta": 0.01}))
    shock("Target and exit multiples", "-10%", "+10%",
          lambda: _pt(a, {"multiple_delta": -0.10}), lambda: _pt(a, {"multiple_delta": 0.10}))
    shock("WACC", "+0.5 pt", "-0.5 pt", with_wacc(0.005), with_wacc(-0.005))
    shock("CapEx % revenue", "+1 pt", "-1 pt", with_margins("capex_pct_rev", 0.01),
          with_margins("capex_pct_rev", -0.01))
    shock("Tax rate", "+1 pt", "-1 pt", with_margins("tax_rate", 0.01), with_margins("tax_rate", -0.01))
    for d in out:
        d["base"] = base
    return sorted(out, key=lambda d: d["swing"], reverse=True)


def write_to_workbook(path, mc: dict, drv: list[dict]) -> None:
    wb = load_workbook(path)
    for name in ("Monte Carlo", "Value Drivers"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("Monte Carlo")
    ws["A1"], ws["A1"].font = "Monte Carlo: base-case price target distribution", TITLE
    ws["A2"] = (f"Script output (values, not formulas): {mc['n']:,} runs of the model engine, seed "
                f"{mc['seed']}. Growth, margin, multiple and WACC deltas are drawn from normal "
                "distributions with the bull and bear cases at about the 90th and 10th percentiles.")
    ws["A2"].font = NOTE
    r = 4
    ws.cell(row=r, column=1, value="Assumption").font = BOLD
    ws.cell(row=r, column=2, value="Std dev").font = BOLD
    for k, v in mc["sigmas"].items():
        r += 1
        ws.cell(row=r, column=1, value=k.replace("_", " "))
        ws.cell(row=r, column=2, value=v).number_format = "0.00%"
    r += 2
    ws.cell(row=r, column=1, value="Result").font = BOLD
    rows = [("Current price", mc["price"], "$#,##0.00"), ("Mean price target", mc["mean"], "$#,##0.00")]
    rows += [(f"P{p}", v, "$#,##0.00") for p, v in mc["percentiles"].items()]
    rows += [("Chance the target is above today's price", mc["p_above_price"], "0%"),
             ("Chance the total return beats the S&P 500 benchmark", mc["p_beats_benchmark"], "0%")]
    for label, v, fmt in rows:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=v).number_format = fmt
    r += 2
    hist_top = r
    ws.cell(row=r, column=1, value="Price target bucket (from)").font = BOLD
    ws.cell(row=r, column=2, value="Runs").font = BOLD
    for lo, cnt in zip(mc["histogram"]["edges"][:-1], mc["histogram"]["counts"], strict=True):
        r += 1
        ws.cell(row=r, column=1, value=round(lo, 2)).number_format = "$#,##0"
        ws.cell(row=r, column=2, value=cnt)
    chart = BarChart()
    chart.title, chart.y_axis.title, chart.x_axis.title = "Simulated price targets", "Runs", "Price target"
    chart.add_data(Reference(ws, min_col=2, min_row=hist_top, max_row=r), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=hist_top + 1, max_row=r))
    chart.legend = None
    chart.width, chart.height = 22, 10
    ws.add_chart(chart, "D4")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 14

    wd = wb.create_sheet("Value Drivers")
    wd["A1"], wd["A1"].font = "Value drivers: base price target, one input at a time", TITLE
    wd["A2"] = "Script output (values). Largest swing first: these are the assumptions the thesis leans on."
    wd["A2"].font = NOTE
    heads = ["Driver", "Low case", "PT (low)", "High case", "PT (high)", "Swing", "Base PT"]
    for j, hd in enumerate(heads):
        c = wd.cell(row=4, column=1 + j, value=hd)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for i, d in enumerate(drv):
        r = 5 + i
        vals = [d["driver"], d["low_case"], d["pt_low"], d["high_case"], d["pt_high"], d["swing"], d["base"]]
        for j, v in enumerate(vals):
            c = wd.cell(row=r, column=1 + j, value=v)
            if j in (2, 4, 5, 6):
                c.number_format = "$#,##0.00"
    chart = BarChart()
    chart.type, chart.title = "bar", "Price target swing by driver"
    chart.add_data(Reference(wd, min_col=6, min_row=4, max_row=4 + len(drv)), titles_from_data=True)
    chart.set_categories(Reference(wd, min_col=1, min_row=5, max_row=4 + len(drv)))
    chart.legend = None
    chart.width, chart.height = 20, 9
    wd.add_chart(chart, "A13")
    wd.column_dimensions["A"].width = 32
    for col in "BCDEFG":
        wd.column_dimensions[col].width = 14
    wb.save(path)
